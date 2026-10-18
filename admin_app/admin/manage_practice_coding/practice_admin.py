import logging
import os
import re
from datetime import datetime

import mysql.connector
from flask import jsonify, render_template, request, session
from werkzeug.utils import secure_filename

from admin import admin_bp
from admin.upload_utils import is_real_image
from config import Config
from extensions import admin_login_required

logger = logging.getLogger(__name__)

from admin.shared_files import publish as publish_upload
UPLOAD_FOLDER = os.path.join(Config.UPLOAD_BASE, 'practice_coding')
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
ALLOWED_IMAGE_EXT = {'png', 'jpg', 'jpeg', 'gif', 'webp'}
IMAGE_SLOTS = ('statement_image', 'input_image', 'example_image')
LANGUAGES = ('python', 'java', 'c', 'cpp', 'csharp', 'javascript')
DIFFICULTIES = ('Easy', 'Moderate', 'Hard')
SUBJECT_SUGGESTIONS = ('String', 'Array', 'Hash Table', 'Math', 'Sorting',
                       'Dynamic Programming', 'Tree', 'Graph', 'Stack',
                       'Queue', 'Linked List', 'Recursion', 'General')


class PracticeDB:
    def _conn(self):
        from admin.db_pool import get_connection
        return get_connection()

    def all(self, sql, params=None):
        conn = self._conn()
        cur = conn.cursor(dictionary=True)
        try:
            cur.execute(sql, params or ())
            return cur.fetchall()
        finally:
            cur.close()
            conn.close()

    def one(self, sql, params=None):
        rows = self.all(sql, params)
        return rows[0] if rows else None

    def write(self, sql, params=None):
        conn = self._conn()
        cur = conn.cursor()
        try:
            cur.execute(sql, params or ())
            conn.commit()
            return cur.lastrowid or cur.rowcount
        except mysql.connector.Error:
            conn.rollback()
            raise
        finally:
            cur.close()
            conn.close()

    def writemany(self, sql, rows):
        if not rows:
            return 0
        conn = self._conn()
        cur = conn.cursor()
        try:
            cur.executemany(sql, rows)
            conn.commit()
            return cur.rowcount
        except mysql.connector.Error:
            conn.rollback()
            raise
        finally:
            cur.close()
            conn.close()


db = PracticeDB()


def _slugify(text):
    slug = re.sub(r'[^a-z0-9]+', '-', (text or '').lower()).strip('-')
    return slug[:88]


def _unique_slug(title, question_id=None):
    base = _slugify(title) or 'problem'
    slug = base
    n = 2
    while True:
        if question_id:
            clash = db.one("SELECT id FROM pcp_questions WHERE slug = %s AND id != %s",
                           (slug, question_id))
        else:
            clash = db.one("SELECT id FROM pcp_questions WHERE slug = %s", (slug,))
        if not clash:
            return slug
        slug = f"{base}-{n}"
        n += 1


def _link_companies(question_id, names):
    db.write("DELETE FROM pcp_question_companies WHERE question_id = %s", (question_id,))
    seen = set()
    for raw in names or []:
        name = (raw or '').strip()[:80]
        if not name or name.lower() in seen:
            continue
        seen.add(name.lower())
        row = db.one("SELECT id FROM pcp_companies WHERE name = %s", (name,))
        company_id = row['id'] if row else db.write(
            "INSERT INTO pcp_companies (name, slug) VALUES (%s, %s)",
            (name, _slugify(name) or 'company'))
        db.write("INSERT IGNORE INTO pcp_question_companies (question_id, company_id) "
                 "VALUES (%s, %s)", (question_id, company_id))


def _companies_of(question_id):
    return [r['name'] for r in db.all(
        "SELECT c.name FROM pcp_question_companies qc "
        "JOIN pcp_companies c ON c.id = qc.company_id "
        "WHERE qc.question_id = %s ORDER BY c.name", (question_id,))]


def _err(message, status=400):
    return jsonify({'success': False, 'message': message}), status


def _text(value, limit=None):
    text = value if isinstance(value, str) else ('' if value is None else str(value))
    text = text.strip()
    if limit and len(text) > limit:
        raise ValueError(f'Value is too long (limit {limit} characters).')
    return text or None


def _int(value, default, low, high, label):
    try:
        n = int(value) if value not in (None, '') else default
    except (TypeError, ValueError):
        raise ValueError(f'{label} must be a whole number.')
    if n < low or n > high:
        raise ValueError(f'{label} must be between {low} and {high}.')
    return n


COMPARE_MODES = ('exact', 'float', 'unordered')
MEMORY_RE = re.compile(r'^(\d+)\s*([kmg]?)b?$', re.I)
MEMORY_MIN_MB, MEMORY_MAX_MB = 16, 4096
DIFFICULTY_ALIASES = {'medium': 'Moderate', 'moderate': 'Moderate', 'easy': 'Easy', 'hard': 'Hard'}


def _memory_limit(value):
    text = str(value if value is not None else '').strip()
    if not text:
        return '512m'
    m = MEMORY_RE.match(text)
    if not m:
        raise ValueError('Memory limit must be a number with a unit, for example 256m, 512m or 1g.')
    amount, unit = int(m.group(1)), (m.group(2) or 'm').lower()
    mb = {'k': amount / 1024, 'm': amount, 'g': amount * 1024}[unit]
    if mb < MEMORY_MIN_MB or mb > MEMORY_MAX_MB:
        raise ValueError(f'Memory limit must be between {MEMORY_MIN_MB}m and {MEMORY_MAX_MB // 1024}g.')
    return f'{amount}{unit}'


def _plain(value):
    text = re.sub(r'<[^>]+>', ' ', str(value or ''))
    return re.sub(r'\s+', ' ', text.replace('\xa0', ' ')).strip().lower()


def question_fingerprint(title, statement):
    return f'{_plain(title)}|{_plain(statement)}'


def unescape_cell(value):
    text = (value or '').strip()
    return text.replace('\\\\n', '\x00').replace('\\n', '\n').replace('\x00', '\\n')


def _compare_settings(data):
    mode = (str(data.get('compare_mode') or 'exact')).strip().lower()
    if mode not in COMPARE_MODES:
        raise ValueError("Comparison must be exact, float, or unordered.")
    tolerance = None
    if mode == 'float':
        raw = data.get('float_tolerance')
        if raw in (None, ''):
            tolerance = 1e-6
        else:
            try:
                tolerance = float(raw)
            except (TypeError, ValueError):
                raise ValueError('Float tolerance must be a number, e.g. 0.000001.')
            if tolerance <= 0:
                raise ValueError('Float tolerance must be greater than zero.')
    return mode, tolerance


@admin_bp.route('/manage-practice-coding')
@admin_login_required
def manage_practice_coding():
    stats = db.one("""
        SELECT COUNT(*) AS total,
               SUM(is_paid = 1) AS paid,
               SUM(is_paid = 0) AS free
        FROM pcp_questions WHERE is_active = 1
    """) or {'total': 0, 'paid': 0, 'free': 0}
    return render_template('admin/manage_practice_coding/practice_coding.html',
                           stats=stats,
                           languages=LANGUAGES,
                           difficulties=DIFFICULTIES,
                           subjects=SUBJECT_SUGGESTIONS)


@admin_bp.route('/api/practice-coding/questions', methods=['GET'])
@admin_login_required
def pcp_list_questions():
    active = 0 if request.args.get('status') == 'removed' else 1
    sql = """
        SELECT q.id, q.title, q.difficulty, q.subject, q.is_paid, q.is_active,
               q.statement_image, q.input_image, q.example_image,
               (SELECT COUNT(*) FROM pcp_testcases t WHERE t.question_id = q.id) AS testcase_count,
               (SELECT COUNT(*) FROM pcp_submissions s WHERE s.question_id = q.id) AS attempt_count
        FROM pcp_questions q WHERE q.is_active = %s
    """
    params = [active]
    subject = request.args.get('subject')
    if subject:
        sql += " AND q.subject = %s"
        params.append(subject)
    tier = request.args.get('tier')
    if tier in ('free', 'paid'):
        sql += " AND q.is_paid = %s"
        params.append(1 if tier == 'paid' else 0)
    sql += " ORDER BY q.position, q.id"
    questions = db.all(sql, tuple(params))
    for q in questions:
        q['companies'] = _companies_of(q['id'])
    return jsonify({'success': True, 'questions': questions,
                    'subjects': db.all("""SELECT subject, COUNT(*) AS n FROM pcp_questions
                                          WHERE is_active = 1 GROUP BY subject ORDER BY subject""")})


@admin_bp.route('/api/practice-coding/questions/<int:qid>', methods=['GET'])
@admin_login_required
def pcp_get_question(qid):
    question = db.one("SELECT * FROM pcp_questions WHERE id = %s AND is_active = 1", (qid,))
    if not question:
        return _err('Question not found.', 404)
    question['allowed_languages'] = [l for l in (question.get('allowed_languages') or '').split(',') if l]
    question['companies'] = _companies_of(qid)
    return jsonify({
        'success': True, 'question': question,
        'testcases': db.all("SELECT * FROM pcp_testcases WHERE question_id = %s ORDER BY position, id", (qid,)),
        'templates': {t['language']: t['template_code'] for t in
                      db.all("SELECT language, template_code FROM pcp_templates WHERE question_id = %s", (qid,))},
    })


def _build_payload(data):
    title = _text(data.get('title'), 255)
    if not title:
        raise ValueError('Question title is required.')

    statement = _text(data.get('statement_html'))
    statement_image = _text(data.get('statement_image'), 255)
    if not statement and not statement_image:
        raise ValueError('Add the question text, or upload a question image.')

    difficulty = DIFFICULTY_ALIASES.get((data.get('difficulty') or 'Easy').strip().lower())
    if difficulty not in DIFFICULTIES:
        raise ValueError(f"Difficulty must be one of: {', '.join(DIFFICULTIES)} (Medium is accepted as Moderate).")

    subject = _text(data.get('subject'), 60) or 'General'

    langs = [l.strip().lower() for l in (data.get('allowed_languages') or []) if l]
    bad = [l for l in langs if l not in LANGUAGES]
    if bad:
        raise ValueError(f"Unknown language: {', '.join(bad)}.")
    if not langs:
        raise ValueError('Select at least one allowed language.')

    testcases = []
    for i, case in enumerate(data.get('testcases') or [], start=1):
        expected = (case.get('expected_output') or '').strip()
        if not expected:
            raise ValueError(f'Test case {i}: expected output is required.')
        testcases.append((case.get('input_text') or '', expected,
                          1 if case.get('is_hidden') else 0, i))
    if not testcases:
        raise ValueError('Add at least one test case so submissions can be judged.')

    compare_mode, float_tolerance = _compare_settings(data)

    return {
        'title': title, 'slug': _unique_slug(title, data.get('id')),
        'difficulty': difficulty, 'subject': subject,
        'statement_html': statement,
        'input_spec': _text(data.get('input_spec')),
        'example_html': _text(data.get('example_html')),
        'statement_image': statement_image,
        'input_image': _text(data.get('input_image'), 255),
        'example_image': _text(data.get('example_image'), 255),
        'allowed_languages': ','.join(langs),
        'output_spec': _text(data.get('output_spec')),
        'time_limit_sec': _int(data.get('time_limit_sec'), 10, 1, 300, 'Time limit'),
        'memory_limit': _memory_limit(data.get('memory_limit')),
        'compare_mode': compare_mode, 'float_tolerance': float_tolerance,
        'is_paid': 1 if data.get('is_paid') else 0,
        'position': _int(data.get('position'), 0, 0, 100000, 'Position'),
    }, testcases


def stored_practice_fingerprints():
    out = {}
    for row in db.all("SELECT id, title, statement_html, is_active FROM pcp_questions"):
        fp = question_fingerprint(row['title'], row['statement_html'])
        out.setdefault(fp, row)
    return out


def duplicate_message(row):
    where = 'in the Removed list - restore it instead' if not row.get('is_active') else f"as #{row['id']} \"{row['title']}\""
    return f'This exact question (same title and statement) is already in the bank {where}. Nothing was added.'


def _save_children(qid, data, testcases):
    db.write("DELETE FROM pcp_testcases WHERE question_id = %s", (qid,))
    db.writemany("""INSERT INTO pcp_testcases (question_id, input_text, expected_output, is_hidden, position)
                    VALUES (%s,%s,%s,%s,%s)""",
                 [(qid,) + tc for tc in testcases])
    db.write("DELETE FROM pcp_templates WHERE question_id = %s", (qid,))
    rows = [(qid, lang, code) for lang, code in (data.get('templates') or {}).items()
            if lang in LANGUAGES and (code or '').strip()]
    db.writemany("INSERT INTO pcp_templates (question_id, language, template_code) VALUES (%s,%s,%s)", rows)


@admin_bp.route('/api/practice-coding/questions', methods=['POST'])
@admin_login_required
def pcp_add_question():
    data = request.get_json(silent=True) or {}
    try:
        payload, testcases = _build_payload(data)
    except ValueError as e:
        return _err(str(e))
    already = stored_practice_fingerprints().get(question_fingerprint(payload['title'], payload['statement_html']))
    if already:
        return _err(duplicate_message(already), 409)
    try:
        columns = ', '.join(payload)
        marks = ', '.join(['%s'] * len(payload))
        qid = db.write(f"INSERT INTO pcp_questions ({columns}) VALUES ({marks})",
                       tuple(payload.values()))
        _save_children(qid, data, testcases)
        _link_companies(qid, data.get('companies'))
    except Exception:
        logger.exception('PRACTICE ADMIN: add question failed')
        return _err('Could not save the question.', 500)
    logger.info("PRACTICE ADMIN: [%s] added %s question %s (%s)",
                session.get('admin_email'), 'PAID' if payload['is_paid'] else 'FREE',
                qid, payload['title'][:40])
    return jsonify({'success': True, 'question_id': qid})


@admin_bp.route('/api/practice-coding/questions/<int:qid>', methods=['PUT'])
@admin_login_required
def pcp_update_question(qid):
    if not db.one("SELECT id FROM pcp_questions WHERE id = %s AND is_active = 1", (qid,)):
        return _err('Question not found.', 404)
    data = request.get_json(silent=True) or {}
    data['id'] = qid
    try:
        payload, testcases = _build_payload(data)
    except ValueError as e:
        return _err(str(e))
    payload.pop('slug', None)
    already = stored_practice_fingerprints().get(question_fingerprint(payload['title'], payload['statement_html']))
    if already and already['id'] != qid:
        return _err(duplicate_message(already), 409)
    try:
        assigns = ', '.join(f"{c} = %s" for c in payload)
        db.write(f"UPDATE pcp_questions SET {assigns} WHERE id = %s",
                 tuple(payload.values()) + (qid,))
        _save_children(qid, data, testcases)
        _link_companies(qid, data.get('companies'))
    except Exception:
        logger.exception('PRACTICE ADMIN: update question %s failed', qid)
        return _err('Could not update the question.', 500)
    return jsonify({'success': True})


@admin_bp.route('/api/practice-coding/questions/<int:qid>', methods=['DELETE'])
@admin_login_required
def pcp_delete_question(qid):
    if not db.one("SELECT id FROM pcp_questions WHERE id = %s", (qid,)):
        return _err('Question not found.', 404)
    try:
        db.write("UPDATE pcp_questions SET is_active = 0 WHERE id = %s", (qid,))
    except Exception:
        logger.exception('PRACTICE ADMIN: delete question %s failed', qid)
        return _err('Could not delete the question.', 500)
    return jsonify({'success': True})


@admin_bp.route('/api/practice-coding/questions/<int:qid>/restore', methods=['PUT'])
@admin_login_required
def pcp_restore_question(qid):
    row = db.one("SELECT is_active FROM pcp_questions WHERE id = %s", (qid,))
    if not row:
        return _err('Question not found.', 404)
    if row['is_active']:
        return jsonify({'success': True, 'already': True})
    db.write("UPDATE pcp_questions SET is_active = 1 WHERE id = %s", (qid,))
    logger.info("PRACTICE ADMIN: [%s] restored question %s", session.get('admin_email'), qid)
    return jsonify({'success': True})


@admin_bp.route('/api/practice-coding/questions/<int:qid>/tier', methods=['PUT'])
@admin_login_required
def pcp_toggle_tier(qid):
    row = db.one("SELECT is_paid FROM pcp_questions WHERE id = %s AND is_active = 1", (qid,))
    if not row:
        return _err('Question not found.', 404)
    new_value = 0 if row['is_paid'] else 1
    db.write("UPDATE pcp_questions SET is_paid = %s WHERE id = %s", (new_value, qid))
    logger.info("PRACTICE ADMIN: [%s] set question %s to %s",
                session.get('admin_email'), qid, 'PAID' if new_value else 'FREE')
    return jsonify({'success': True, 'is_paid': new_value})


@admin_bp.route('/api/practice-coding/image', methods=['POST'])
@admin_login_required
def pcp_upload_image():
    slot = (request.form.get('slot') or 'statement_image').strip()
    if slot not in IMAGE_SLOTS:
        return _err('Unknown image slot.')
    file = request.files.get('image')
    if not file or not file.filename:
        return _err('No image was selected.')
    ext = file.filename.rsplit('.', 1)[-1].lower() if '.' in file.filename else ''
    if ext not in ALLOWED_IMAGE_EXT:
        return _err('Only PNG, JPG, GIF or WEBP images are allowed.')
    if not is_real_image(file):
        return _err('That file is not a valid image.')
    name = f"{datetime.now().strftime('%Y%m%d%H%M%S%f')}_{slot}_{secure_filename(file.filename)}"
    try:
        file.save(os.path.join(UPLOAD_FOLDER, name))
        publish_upload(os.path.join(UPLOAD_FOLDER, name))
    except OSError:
        logger.exception('PRACTICE ADMIN: image save failed')
        return _err('Could not save the image.', 500)
    return jsonify({'success': True, 'slot': slot,
                    'url': '/static/' + os.path.join('uploads', 'practice_coding', name).replace('\\', '/')})
BULK_COLUMNS = ['title', 'difficulty', 'subject', 'access', 'companies',
                'compare_mode', 'float_tolerance',
                'statement', 'statement_image',
                'input_format', 'input_image',
                'example', 'example_image',
                'output_format', 'languages',
                'template_python', 'template_java', 'template_c',
                'template_cpp', 'template_csharp', 'template_javascript',
                'tc1_input', 'tc1_output', 'tc2_input', 'tc2_output',
                'tc3_input', 'tc3_output',
                'hidden1_input', 'hidden1_output', 'hidden2_input', 'hidden2_output']

BULK_MAX_ROWS = 500

BULK_SAMPLE_ROW = [
    'Two Sum', 'Easy', 'Array', 'free', 'Google; Amazon',
    'exact', '',
    'Given n integers and a target, print the indices of the two numbers that add up to it.',
    '',
    'First line: n and target. Second line: n space-separated integers.',
    '',
    'Input:\\n4 9\\n2 7 11 15\\n\\nOutput:\\n0 1',
    '',
    'Two space-separated indices.',
    'python,java',
    '# read input here\\nprint()', 'public class Main { }', '', '', '', '',
    '4 9\\n2 7 11 15', '0 1', '', '', '', '',
    '3 6\\n3 3 1', '0 1', '', '']


def _cell(value):
    return unescape_cell(value)


def _read_upload_rows(file):
    name = (file.filename or '').lower()
    if name.endswith('.xlsx'):
        try:
            from openpyxl import load_workbook
        except ImportError:
            return None, 'Excel support is not installed on the server. Use the CSV template instead.'
        try:
            wb = load_workbook(file, read_only=True, data_only=True)
        except Exception:
            return None, 'Could not read that Excel file. Re-save it as .xlsx, or use the CSV template.'
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        header = next(it, None)
        if not header:
            return [], None
        header = [(str(h).strip() if h is not None else '') for h in header]
        rows = []
        for r in it:
            rows.append({k: ('' if v is None else str(v))
                         for k, v in zip(header, r) if k})
        return rows, None
    if name.endswith('.csv'):
        import csv
        from io import StringIO
        try:
            text = file.read().decode('utf-8-sig')
        except UnicodeDecodeError:
            return None, 'The file is not UTF-8. In Excel, save as "CSV UTF-8".'
        return list(csv.DictReader(StringIO(text))), None
    return None, 'Upload a .csv or .xlsx file (use the template above).'


def _testcases_from_row(row):
    found = {}
    for key in row.keys():
        if not key:
            continue
        m = re.match(r'^(tc|hidden)(\d+)_(input|output)$', str(key).strip().lower())
        if m:
            found.setdefault((m.group(1), int(m.group(2))), True)
    cases = []
    for kind, n in sorted(found.keys(), key=lambda k: (k[0] != 'tc', k[1])):
        out_val = _cell(row.get(f'{kind}{n}_output'))
        in_val = _cell(row.get(f'{kind}{n}_input'))
        if out_val:
            cases.append({'input_text': in_val, 'expected_output': out_val,
                          'is_hidden': kind == 'hidden'})
    return cases


def _templates_from_row(row):
    out = {}
    for lang in LANGUAGES:
        code = row.get('template_' + lang)
        if code and str(code).strip():
            out[lang] = str(code).replace('\\n', '\n')
    return out


def _bump_slug_in_batch(slug, taken):
    if slug not in taken:
        return slug
    n = 2
    while (f'{slug}-{n}' in taken
           or db.one("SELECT id FROM pcp_questions WHERE slug = %s", (f'{slug}-{n}',))):
        n += 1
    return f'{slug}-{n}'


@admin_bp.route('/api/practice-coding/bulk/template', methods=['GET'])
@admin_login_required
def pcp_bulk_template():
    fmt = (request.args.get('fmt') or 'csv').lower()
    if fmt == 'xlsx':
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            return _err('Excel support is not installed. Use the CSV template.', 501)
        from io import BytesIO
        from flask import send_file
        wb = Workbook()
        ws = wb.active
        ws.title = 'Questions'
        ws.append(BULK_COLUMNS)
        ws.append([c.replace('\\n', '\n') for c in BULK_SAMPLE_ROW])
        fill = PatternFill('solid', fgColor='833D17')
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = fill
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, as_attachment=True,
                         download_name='practice_coding_template.xlsx',
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    import csv
    from io import StringIO
    from flask import Response
    out = StringIO()
    writer = csv.writer(out)
    writer.writerow(BULK_COLUMNS)
    writer.writerow(BULK_SAMPLE_ROW)
    return Response(out.getvalue(), mimetype='text/csv',
                    headers={'Content-Disposition':
                             'attachment; filename=practice_coding_template.csv'})


def _bulk_insert_transactional(prepared):
    conn = db._conn()
    conn.autocommit = False
    cur = conn.cursor()
    try:
        for payload, data, cases, companies in prepared:
            columns = ', '.join(payload)
            marks = ', '.join(['%s'] * len(payload))
            cur.execute(f"INSERT INTO pcp_questions ({columns}) VALUES ({marks})",
                        tuple(payload.values()))
            qid = cur.lastrowid
            if cases:
                cur.executemany(
                    "INSERT INTO pcp_testcases (question_id, input_text, expected_output, "
                    "is_hidden, position) VALUES (%s,%s,%s,%s,%s)",
                    [(qid,) + tc for tc in cases])
            tpl_rows = [(qid, lang, code) for lang, code in (data.get('templates') or {}).items()
                        if lang in LANGUAGES and (code or '').strip()]
            if tpl_rows:
                cur.executemany(
                    "INSERT INTO pcp_templates (question_id, language, template_code) "
                    "VALUES (%s,%s,%s)", tpl_rows)
            seen = set()
            for raw in companies or []:
                cname = (raw or '').strip()[:80]
                if not cname or cname.lower() in seen:
                    continue
                seen.add(cname.lower())
                cur.execute("SELECT id FROM pcp_companies WHERE name = %s", (cname,))
                r = cur.fetchone()
                if r:
                    company_id = r[0]
                else:
                    cur.execute("INSERT INTO pcp_companies (name, slug) VALUES (%s,%s)",
                                (cname, _slugify(cname) or 'company'))
                    company_id = cur.lastrowid
                cur.execute("INSERT IGNORE INTO pcp_question_companies "
                            "(question_id, company_id) VALUES (%s,%s)", (qid, company_id))
        conn.commit()
        return len(prepared)
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/practice-coding/bulk/upload', methods=['POST'])
@admin_login_required
def pcp_bulk_upload():
    file = request.files.get('file')
    if not file or not file.filename:
        return _err('Choose a CSV or Excel file to upload.')

    rows, read_error = _read_upload_rows(file)
    if read_error:
        return _err(read_error)
    if not rows:
        return _err('The file has no data rows.')
    if len(rows) > BULK_MAX_ROWS:
        return _err(f'Maximum {BULK_MAX_ROWS} questions per upload.')

    prepared, errors, batch_slugs, skipped, seen = [], [], set(), [], {}
    stored = stored_practice_fingerprints()
    for i, row in enumerate(rows, start=2):
        if not any(str(v or '').strip() for v in row.values()):
            continue
        try:
            access = (row.get('access') or 'free').strip().lower()
            if access not in ('free', 'paid'):
                raise ValueError("access must be 'free' or 'paid'")
            languages = [l.strip().lower() for l in
                         (row.get('languages') or 'python').split(',') if l.strip()]
            data = {
                'title': row.get('title'), 'difficulty': row.get('difficulty') or 'Easy',
                'subject': row.get('subject') or 'General', 'is_paid': access == 'paid',
                'statement_html': _cell(row.get('statement')),
                'statement_image': _cell(row.get('statement_image')),
                'input_spec': _cell(row.get('input_format')),
                'input_image': _cell(row.get('input_image')),
                'example_html': _cell(row.get('example')),
                'example_image': _cell(row.get('example_image')),
                'output_spec': _cell(row.get('output_format')),
                'allowed_languages': languages,
                'compare_mode': row.get('compare_mode') or 'exact',
                'float_tolerance': row.get('float_tolerance'),
                'templates': _templates_from_row(row),
                'testcases': _testcases_from_row(row),
            }
            payload, cases = _build_payload(data)
            fp = question_fingerprint(payload['title'], payload['statement_html'])
            if fp in stored:
                skipped.append(f"Row {i}: already in the bank as #{stored[fp]['id']} - skipped.")
                continue
            if fp in seen:
                skipped.append(f'Row {i}: same as row {seen[fp]} of this file - skipped.')
                continue
            seen[fp] = i
            payload['slug'] = _bump_slug_in_batch(payload['slug'], batch_slugs)
            batch_slugs.add(payload['slug'])
            companies = [c for c in (row.get('companies') or '').replace(';', ',').split(',')
                         if c.strip()]
            prepared.append((payload, data, cases, companies))
        except ValueError as e:
            errors.append(f'Row {i}: {e}')

    if errors:
        return jsonify({'success': False,
                        'message': f'{len(errors)} row(s) need fixing. Nothing was imported.',
                        'errors': errors[:50]}), 400
    if not prepared:
        if skipped:
            return jsonify({'success': True, 'imported': 0, 'skipped': skipped,
                            'message': f'Nothing added - all {len(skipped)} question(s) in this file are already in the bank.'})
        return _err('The file has no fillable rows.')

    try:
        imported = _bulk_insert_transactional(prepared)
    except Exception:
        logger.exception('PRACTICE ADMIN: bulk import transaction failed')
        return _err('A database error rolled the whole import back. Nothing was saved.', 500)

    logger.info("PRACTICE ADMIN: [%s] bulk-imported %d coding question(s), %d skipped",
                session.get('admin_email'), imported, len(skipped))
    return jsonify({'success': True, 'imported': imported, 'skipped': skipped,
                    'message': f'Imported {imported} question(s).' + (f' {len(skipped)} already in the bank were skipped.' if skipped else '')})


@admin_bp.route('/api/practice-coding/questions/<int:qid>/editorials', methods=['GET'])
@admin_login_required
def pcp_list_editorials(qid):
    rows = db.all(
        "SELECT id, position, title, intro_html, outro_html, time_complexity, "
        "space_complexity, is_published FROM pcp_editorials "
        "WHERE question_id = %s ORDER BY position, id", (qid,))
    if rows:
        ids = [r['id'] for r in rows]
        marks = ','.join(['%s'] * len(ids))
        codes = db.all(
            f"SELECT editorial_id, language, code FROM pcp_editorial_code "
            f"WHERE editorial_id IN ({marks})", tuple(ids))
        grouped = {}
        for c in codes:
            grouped.setdefault(c['editorial_id'], {})[c['language']] = c['code']
        for r in rows:
            r['codes'] = grouped.get(r['id'], {})
    return jsonify({'editorials': rows})


def _editorial_payload(data):
    title = _text(data.get('title'), 255)
    if not title:
        return None, 'An approach needs a title, for example "Brute Force".'

    codes = data.get('codes') or {}
    if not isinstance(codes, dict):
        return None, 'Solution code must be sent as an object keyed by language.'
    clean_codes = {lang: str(code) for lang, code in codes.items()
                   if lang in LANGUAGES and str(code or '').strip()}

    return {
        'title': title,
        'intro_html': _text(data.get('intro_html')) or None,
        'outro_html': _text(data.get('outro_html')) or None,
        'time_complexity': _text(data.get('time_complexity'), 120) or None,
        'space_complexity': _text(data.get('space_complexity'), 120) or None,
        'position': _int(data.get('position'), 0, 0, 99, 'Position'),
        'is_published': 0 if str(data.get('is_published', '1')) in ('0', 'false', 'False') else 1,
        'codes': clean_codes,
    }, None


def _save_editorial_codes(editorial_id, codes):
    db.write("DELETE FROM pcp_editorial_code WHERE editorial_id = %s", (editorial_id,))
    rows = [(editorial_id, lang, code) for lang, code in codes.items()]
    if rows:
        db.writemany(
            "INSERT INTO pcp_editorial_code (editorial_id, language, code) "
            "VALUES (%s, %s, %s)", rows)


@admin_bp.route('/api/practice-coding/questions/<int:qid>/editorials', methods=['POST'])
@admin_login_required
def pcp_add_editorial(qid):
    if not db.one("SELECT id FROM pcp_questions WHERE id = %s", (qid,)):
        return _err('That question no longer exists.', 404)

    payload, error = _editorial_payload(request.get_json(silent=True) or {})
    if error:
        return _err(error)

    try:
        eid = db.write(
            "INSERT INTO pcp_editorials (question_id, position, title, intro_html, "
            "outro_html, time_complexity, space_complexity, is_published) "
            "VALUES (%s,%s,%s,%s,%s,%s,%s,%s)",
            (qid, payload['position'], payload['title'], payload['intro_html'],
             payload['outro_html'], payload['time_complexity'],
             payload['space_complexity'], payload['is_published']))
        _save_editorial_codes(eid, payload['codes'])
    except mysql.connector.Error:
        logger.exception("PRACTICE ADMIN: could not add editorial for question %s", qid)
        return _err('Could not save the approach.', 500)

    logger.info("PRACTICE ADMIN: [%s] added editorial %s to question %s",
                session.get('admin_email'), eid, qid)
    return jsonify({'success': True, 'id': eid}), 201


@admin_bp.route('/api/practice-coding/editorials/<int:eid>', methods=['PUT'])
@admin_login_required
def pcp_update_editorial(eid):
    if not db.one("SELECT id FROM pcp_editorials WHERE id = %s", (eid,)):
        return _err('That approach no longer exists.', 404)

    payload, error = _editorial_payload(request.get_json(silent=True) or {})
    if error:
        return _err(error)

    try:
        db.write(
            "UPDATE pcp_editorials SET position=%s, title=%s, intro_html=%s, "
            "outro_html=%s, time_complexity=%s, space_complexity=%s, is_published=%s "
            "WHERE id = %s",
            (payload['position'], payload['title'], payload['intro_html'],
             payload['outro_html'], payload['time_complexity'],
             payload['space_complexity'], payload['is_published'], eid))
        _save_editorial_codes(eid, payload['codes'])
    except mysql.connector.Error:
        logger.exception("PRACTICE ADMIN: could not update editorial %s", eid)
        return _err('Could not update the approach.', 500)

    logger.info("PRACTICE ADMIN: [%s] updated editorial %s",
                session.get('admin_email'), eid)
    return jsonify({'success': True})


@admin_bp.route('/api/practice-coding/editorials/<int:eid>', methods=['DELETE'])
@admin_login_required
def pcp_delete_editorial(eid):
    removed = db.write("DELETE FROM pcp_editorials WHERE id = %s", (eid,))
    if not removed:
        return _err('That approach no longer exists.', 404)
    logger.warning("PRACTICE ADMIN: [%s] deleted editorial %s",
                   session.get('admin_email'), eid)
    return jsonify({'success': True})
