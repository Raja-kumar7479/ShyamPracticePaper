import csv
import logging
import os
from datetime import datetime
from io import StringIO

from flask import Response, jsonify, render_template, request, session
from werkzeug.utils import secure_filename

from admin import admin_bp
from admin.upload_utils import is_real_image
from admin.manage_practice_coding.practice_admin import (db, _err, _text, _int, _plain, unescape_cell,
                                                         _read_upload_rows)
from config import Config
from extensions import admin_login_required

logger = logging.getLogger(__name__)

from admin.shared_files import publish as publish_upload
QUIZ_UPLOAD_FOLDER = os.path.join(Config.UPLOAD_BASE, 'practice_quiz')
os.makedirs(QUIZ_UPLOAD_FOLDER, exist_ok=True)
ALLOWED_IMAGE_EXT = {'png', 'jpg', 'jpeg', 'gif', 'webp'}
OPTIONS = ('A', 'B', 'C', 'D')
QUIZ_SUBJECTS = ('Java', 'C', 'C++', 'Python', 'Algorithms', 'Data Structures',
                 'SQL', 'DBMS', 'Operating Systems', 'Networking', 'Aptitude')


@admin_bp.route('/manage-practice-quiz')
@admin_login_required
def manage_practice_quiz():
    stats = db.one("""
        SELECT COUNT(*) AS total, SUM(is_paid = 1) AS paid, SUM(is_paid = 0) AS free
        FROM pcq_questions WHERE is_active = 1
    """) or {'total': 0, 'paid': 0, 'free': 0}
    return render_template('admin/manage_practice_coding/practice_quiz.html',
                           stats=stats, subjects=QUIZ_SUBJECTS)


def _build_quiz_payload(data):
    subject = _text(data.get('subject'), 60)
    if not subject:
        raise ValueError('Subject is required - it is how students pick a quiz.')
    if '/' in subject or '\\' in subject:
        raise ValueError('Subject cannot contain slashes. Use "C-CPP" style names instead.')

    question_text = _text(data.get('question_text'))
    question_image = _text(data.get('question_image'), 255)
    if not question_text and not question_image:
        raise ValueError('Type the question, or upload a question image.')

    options = {}
    for letter in OPTIONS:
        value = _text(data.get('option_' + letter.lower()))
        if not value:
            raise ValueError(f'Option {letter} is required - MCQs need all four options.')
        options[letter] = value

    correct = (data.get('correct_option') or '').strip().upper()
    if correct not in OPTIONS:
        raise ValueError('Mark which option (A-D) is the correct answer.')

    return {
        'subject': subject,
        'question_text': question_text,
        'question_image': question_image,
        'option_a': options['A'], 'option_b': options['B'],
        'option_c': options['C'], 'option_d': options['D'],
        'correct_option': correct,
        'explanation': _text(data.get('explanation')),
        'is_paid': 1 if data.get('is_paid') else 0,
        'position': _int(data.get('position'), 0, 0, 100000, 'Position'),
    }


def quiz_fingerprint(p):
    return '|'.join([_plain(p.get('question_text')), _plain(p.get('option_a')), _plain(p.get('option_b')),
                     _plain(p.get('option_c')), _plain(p.get('option_d'))])


def stored_quiz_fingerprints():
    out = {}
    for row in db.all("SELECT id, subject, question_text, question_image, option_a, option_b, option_c, option_d, is_active FROM pcq_questions"):
        if not row.get('question_text'):
            continue
        out.setdefault(quiz_fingerprint(row), row)
    return out


def quiz_duplicate_message(row):
    where = 'in the Removed list - restore it instead' if not row.get('is_active') else f"as #{row['id']} under {row['subject']}"
    return f'This exact question (same text and options) is already in the bank {where}. Nothing was added.'


@admin_bp.route('/api/practice-quiz/questions', methods=['GET'])
@admin_login_required
def pcq_quiz_list():
    active = 0 if request.args.get('status') == 'removed' else 1
    sql = """
        SELECT q.id, q.subject, q.question_text, q.question_image, q.correct_option,
               q.is_paid, q.is_active,
               (SELECT COUNT(*) FROM pcq_attempt_answers a WHERE a.question_id = q.id) AS answer_count
        FROM pcq_questions q WHERE q.is_active = %s
    """
    params = [active]
    subject = request.args.get('subject')
    if subject:
        sql += " AND q.subject = %s"
        params.append(subject)
    tier = request.args.get('tier')
    if tier in ('free', 'paid'):
        sql += " AND q.is_paid = %s"
        params.append(1 if tier == 'paid' else 0)
    sql += " ORDER BY q.subject, q.position, q.id"
    return jsonify({'success': True, 'questions': db.all(sql, tuple(params)),
                    'subjects': db.all("""SELECT subject, COUNT(*) AS n FROM pcq_questions
                                          WHERE is_active = 1 GROUP BY subject ORDER BY subject""")})


@admin_bp.route('/api/practice-quiz/questions/<int:qid>', methods=['GET'])
@admin_login_required
def pcq_quiz_get(qid):
    question = db.one("SELECT * FROM pcq_questions WHERE id = %s AND is_active = 1", (qid,))
    if not question:
        return _err('Question not found.', 404)
    return jsonify({'success': True, 'question': question})


@admin_bp.route('/api/practice-quiz/questions', methods=['POST'])
@admin_login_required
def pcq_quiz_add():
    data = request.get_json(silent=True) or {}
    try:
        payload = _build_quiz_payload(data)
    except ValueError as e:
        return _err(str(e))
    if payload.get('question_text'):
        already = stored_quiz_fingerprints().get(quiz_fingerprint(payload))
        if already:
            return _err(quiz_duplicate_message(already), 409)
    columns = ', '.join(payload)
    marks = ', '.join(['%s'] * len(payload))
    try:
        qid = db.write(f"INSERT INTO pcq_questions ({columns}) VALUES ({marks})",
                       tuple(payload.values()))
    except Exception:
        logger.exception('QUIZ ADMIN: add failed')
        return _err('Could not save the question.', 500)
    logger.info("QUIZ ADMIN: [%s] added %s quiz question %s (%s)",
                session.get('admin_email'), 'PAID' if payload['is_paid'] else 'FREE',
                qid, payload['subject'])
    return jsonify({'success': True, 'question_id': qid})


@admin_bp.route('/api/practice-quiz/questions/<int:qid>', methods=['PUT'])
@admin_login_required
def pcq_quiz_update(qid):
    if not db.one("SELECT id FROM pcq_questions WHERE id = %s AND is_active = 1", (qid,)):
        return _err('Question not found.', 404)
    data = request.get_json(silent=True) or {}
    try:
        payload = _build_quiz_payload(data)
    except ValueError as e:
        return _err(str(e))
    if payload.get('question_text'):
        already = stored_quiz_fingerprints().get(quiz_fingerprint(payload))
        if already and already['id'] != qid:
            return _err(quiz_duplicate_message(already), 409)
    assigns = ', '.join(f"{c} = %s" for c in payload)
    try:
        db.write(f"UPDATE pcq_questions SET {assigns} WHERE id = %s",
                 tuple(payload.values()) + (qid,))
    except Exception:
        logger.exception('QUIZ ADMIN: update %s failed', qid)
        return _err('Could not update the question.', 500)
    return jsonify({'success': True})


@admin_bp.route('/api/practice-quiz/questions/<int:qid>', methods=['DELETE'])
@admin_login_required
def pcq_quiz_delete(qid):
    if not db.one("SELECT id FROM pcq_questions WHERE id = %s", (qid,)):
        return _err('Question not found.', 404)
    db.write("UPDATE pcq_questions SET is_active = 0 WHERE id = %s", (qid,))
    return jsonify({'success': True})


@admin_bp.route('/api/practice-quiz/questions/<int:qid>/restore', methods=['PUT'])
@admin_login_required
def pcq_quiz_restore(qid):
    row = db.one("SELECT is_active FROM pcq_questions WHERE id = %s", (qid,))
    if not row:
        return _err('Question not found.', 404)
    if row['is_active']:
        return jsonify({'success': True, 'already': True})
    db.write("UPDATE pcq_questions SET is_active = 1 WHERE id = %s", (qid,))
    logger.info("QUIZ ADMIN: [%s] restored question %s", session.get('admin_email'), qid)
    return jsonify({'success': True})


@admin_bp.route('/api/practice-quiz/questions/<int:qid>/tier', methods=['PUT'])
@admin_login_required
def pcq_quiz_tier(qid):
    row = db.one("SELECT is_paid FROM pcq_questions WHERE id = %s AND is_active = 1", (qid,))
    if not row:
        return _err('Question not found.', 404)
    new_value = 0 if row['is_paid'] else 1
    db.write("UPDATE pcq_questions SET is_paid = %s WHERE id = %s", (new_value, qid))
    return jsonify({'success': True, 'is_paid': new_value})


@admin_bp.route('/api/practice-quiz/image', methods=['POST'])
@admin_login_required
def pcq_quiz_image():
    file = request.files.get('image')
    if not file or not file.filename:
        return _err('No image was selected.')
    ext = file.filename.rsplit('.', 1)[-1].lower() if '.' in file.filename else ''
    if ext not in ALLOWED_IMAGE_EXT:
        return _err('Only PNG, JPG, GIF or WEBP images are allowed.')
    if not is_real_image(file):
        return _err('That file is not a valid image.')
    name = f"{datetime.now().strftime('%Y%m%d%H%M%S%f')}_quiz_{secure_filename(file.filename)}"
    try:
        file.save(os.path.join(QUIZ_UPLOAD_FOLDER, name))
        publish_upload(os.path.join(QUIZ_UPLOAD_FOLDER, name))
    except OSError:
        logger.exception('QUIZ ADMIN: image save failed')
        return _err('Could not save the image.', 500)
    return jsonify({'success': True,
                    'url': '/static/' + os.path.join('uploads', 'practice_quiz', name).replace('\\', '/')})


QUIZ_BULK_COLUMNS = ['subject', 'question', 'option_a', 'option_b', 'option_c',
                     'option_d', 'correct', 'explanation', 'access']


@admin_bp.route('/api/practice-quiz/bulk/template', methods=['GET'])
@admin_login_required
def pcq_quiz_bulk_template():
    sample = ['Java', 'Which keyword makes a variable constant in Java?',
              'static', 'final', 'const', 'immutable', 'B',
              'final makes the reference unchangeable after assignment.', 'free']
    if (request.args.get('fmt') or 'csv').lower() == 'xlsx':
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            return _err('Excel support is not installed. Use the CSV template.', 501)
        from io import BytesIO
        from flask import send_file
        wb = Workbook()
        ws = wb.active
        ws.title = 'Questions'
        ws.append(QUIZ_BULK_COLUMNS)
        ws.append(sample)
        fill = PatternFill('solid', fgColor='833D17')
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = fill
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name='quiz_questions_template.xlsx',
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    out = StringIO()
    writer = csv.writer(out)
    writer.writerow(QUIZ_BULK_COLUMNS)
    writer.writerow(sample)
    return Response(out.getvalue(), mimetype='text/csv',
                    headers={'Content-Disposition':
                             'attachment; filename=quiz_questions_template.csv'})


@admin_bp.route('/api/practice-quiz/bulk/upload', methods=['POST'])
@admin_login_required
def pcq_quiz_bulk_upload():
    file = request.files.get('file')
    if not file or not file.filename:
        return _err('Choose a CSV or Excel file to upload.')
    rows, read_error = _read_upload_rows(file)
    if read_error:
        return _err(read_error)
    if not rows:
        return _err('The file has no data rows.')
    if len(rows) > 1000:
        return _err('Maximum 1000 questions per upload.')

    prepared, errors, skipped, seen = [], [], [], {}
    stored = stored_quiz_fingerprints()
    for i, row in enumerate(rows, start=2):
        if not any(str(v or '').strip() for v in row.values()):
            continue
        try:
            access = (row.get('access') or 'free').strip().lower()
            if access not in ('free', 'paid'):
                raise ValueError("access must be 'free' or 'paid'")
            payload = _build_quiz_payload({
                'subject': row.get('subject'),
                'question_text': unescape_cell(row.get('question')),
                'option_a': row.get('option_a'), 'option_b': row.get('option_b'),
                'option_c': row.get('option_c'), 'option_d': row.get('option_d'),
                'correct_option': row.get('correct'),
                'explanation': unescape_cell(row.get('explanation')),
                'is_paid': access == 'paid',
            })
            fp = quiz_fingerprint(payload)
            if fp in stored:
                skipped.append(f"Row {i}: already in the bank as #{stored[fp]['id']} - skipped.")
                continue
            if fp in seen:
                skipped.append(f'Row {i}: same as row {seen[fp]} of this file - skipped.')
                continue
            seen[fp] = i
            prepared.append(payload)
        except ValueError as e:
            errors.append(f'Row {i}: {e}')

    if errors:
        return jsonify({'success': False,
                        'message': f'{len(errors)} row(s) need fixing. Nothing was imported.',
                        'errors': errors[:50]}), 400
    if not prepared:
        if skipped:
            return jsonify({'success': True, 'imported': 0, 'skipped': skipped,
                            'message': f'Nothing added - all {len(skipped)} question(s) in this file are already in the bank.'})
        return _err('The file has no fillable rows.')

    conn = db._conn()
    conn.autocommit = False
    cur = conn.cursor()
    try:
        for payload in prepared:
            columns = ', '.join(payload)
            marks = ', '.join(['%s'] * len(payload))
            cur.execute(f"INSERT INTO pcq_questions ({columns}) VALUES ({marks})", tuple(payload.values()))
        conn.commit()
    except Exception:
        conn.rollback()
        logger.exception('QUIZ ADMIN: bulk import rolled back')
        return _err('A database error rolled the whole import back. Nothing was saved.', 500)
    finally:
        cur.close()
        conn.close()

    logger.info("QUIZ ADMIN: [%s] bulk-imported %d quiz question(s), %d skipped",
                session.get('admin_email'), len(prepared), len(skipped))
    return jsonify({'success': True, 'imported': len(prepared), 'skipped': skipped,
                    'message': f'Imported {len(prepared)} question(s).' + (f' {len(skipped)} already in the bank were skipped.' if skipped else '')})
