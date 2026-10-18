import csv
import re
from io import StringIO, BytesIO
import logging
import os
from datetime import datetime
from werkzeug.utils import secure_filename
from flask import render_template, request, jsonify, flash, redirect, url_for, session, Response, send_file
from .practice_paper_db import PaperOperation
from admin.manage_mock_test.test_details_insert import _rows_from_upload, _check_answer, question_fingerprint
from extensions import admin_login_required, admin_owner_required, can_manage
from admin import admin_bp

logger = logging.getLogger(__name__)
paper_op = PaperOperation()

PRACTICE_COLUMNS = ['question_id', 'question_type', 'year', 'paper_set', 'subject', 'topic',
                    'question_text', 'option_a', 'option_b', 'option_c', 'option_d',
                    'correct_option', 'answer_text', 'explanation_link']

PRACTICE_TEMPLATE_ROWS = [
    ['1', 'MCQ', '2024', 'Set A', 'DBMS', 'Normalization',
     '<p>Which normal form removes partial dependency?</p>', '<p>1NF</p>', '<p>2NF</p>', '<p>3NF</p>', '<p>BCNF</p>',
     'B', '<p>2NF removes partial dependencies on a composite key.</p>', ''],
    ['2', 'MSQ', '2024', 'Set A', 'DBMS', 'ACID',
     '<p>Which of these are ACID properties?</p>', '<p>Atomicity</p>', '<p>Availability</p>', '<p>Isolation</p>', '<p>Latency</p>',
     'AC', '<p>Atomicity and Isolation are ACID properties.</p>', ''],
    ['3', 'NAT', '2023', 'Set B', 'Maths', 'Algebra',
     '<p>If 2x = 10, x = ______</p>', '', '', '', '',
     '5', '<p>x = 10 / 2 = 5</p>', ''],
    ['4', 'MCQ', '2023', 'Set B', 'Programming', 'Python',
     '<p>What does this print?</p><pre class="ql-syntax">nums = [1, 2, 3]\nprint(len(nums))</pre>',
     '<p>2</p>', '<p>3</p>', '<p>Error</p>', '<p>None</p>',
     'B', '<p>The list has 3 elements.</p>', ''],
]

PRACTICE_GUIDE = [
    ('question_id', 'Yes', 'A NUMBER unique inside this paper, e.g. 1, 2, 3. Duplicates are rejected.'),
    ('question_type', 'Yes', 'MCQ, MSQ or NAT'),
    ('year', 'No', 'Exam year as a number, e.g. 2024. Leave empty if unknown.'),
    ('paper_set', 'No', 'Paper set name, e.g. Set A'),
    ('subject', 'Yes', 'Subject bucket shown to students, e.g. DBMS'),
    ('topic', 'No', 'Topic inside the subject, e.g. Normalization'),
    ('question_text', 'Yes', 'The question. Plain text or HTML. For code wrap it in <pre class="ql-syntax">...</pre>'),
    ('option_a .. option_d', 'MCQ/MSQ only', 'The four options. Leave empty for NAT'),
    ('correct_option', 'Yes', 'MCQ: one letter (B). MSQ: letters joined (AC). NAT: the numeric answer'),
    ('answer_text', 'No', 'Worked explanation shown with the solution'),
    ('explanation_link', 'No', 'Optional URL to a video/page explanation'),
]
from config import Config
from admin.shared_files import publish as publish_upload
UPLOAD_FOLDER = os.path.join(Config.UPLOAD_BASE, 'images', 'practice_paper')
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
ALLOWED_EXTENSIONS = {'png', 'jpg', 'jpeg', 'gif'}

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


def paper_fingerprint(q):
    return question_fingerprint(q.get('question_type'), q.get('question_text'), q.get('option_a'),
                                q.get('option_b'), q.get('option_c'), q.get('option_d'))


def stored_paper_index(paper_code):
    by_fp, by_qid = {}, {}
    for row in paper_op.content_index(paper_code):
        fp = paper_fingerprint(row)
        if fp:
            by_fp.setdefault(fp, row)
        by_qid[int(row['question_id'])] = (row, fp)
    return by_fp, by_qid


def _validated_answer(data):
    qtype, answer = _check_answer(data.get('question_type'), data.get('correct_option'),
                                  None if str(data.get('question_type') or '').upper() == 'NAT'
                                  else {k: data.get(f'option_{k}') for k in 'abcd'})
    return qtype, answer

def _clean_option_regions(regions):
    if not isinstance(regions, dict) or not regions:
        raise ValueError('Mark at least one option area on the image.')
    clean = {}
    for key in ('a', 'b', 'c', 'd'):
        box = regions.get(key)
        if not isinstance(box, dict):
            continue
        try:
            clean[key] = {k: round(float(box[k]), 2) for k in ('x', 'y', 'w', 'h')}
        except (KeyError, TypeError, ValueError):
            raise ValueError(f'Option {key.upper()} area is invalid. Mark it again.')
    if not clean:
        raise ValueError('Mark at least one option area on the image.')
    return clean


@admin_bp.route('/api/practice-questions/add-image', methods=['POST'])
@admin_login_required
def practice_add_image_question():
    import json as _json
    admin_email = session.get('admin_email', 'UnknownAdmin')
    data = request.get_json(silent=True) or {}
    required = ['paper_code', 'subject', 'question_image', 'option_regions', 'correct_option']
    missing = [f for f in required if not data.get(f)]
    if missing:
        return jsonify({'message': f"Missing: {', '.join(missing)}."}), 400
    if not paper_op.paper_code_exists(data['paper_code']):
        return jsonify({'message': 'Unknown paper code.'}), 400

    try:
        clean = _clean_option_regions(data.get('option_regions'))
    except ValueError as e:
        return jsonify({'message': str(e)}), 400
    answer = str(data['correct_option']).strip().upper()
    for ch in answer:
        if ch.lower() not in clean:
            return jsonify({'message': f'Option {ch} is marked correct but has no area on the image.'}), 400

    qtype = str(data.get('question_type', 'MCQ')).upper()
    if qtype not in ('MCQ', 'MSQ'):
        return jsonify({'message': 'Image questions must be MCQ or MSQ.'}), 400

    question_id = str(data.get('question_id') or '').strip()
    if not question_id:
        question_id = paper_op.next_question_id(data['paper_code'])
    elif paper_op.question_id_exists(data['paper_code'], question_id):
        return jsonify({'message': f'question_id {question_id} already exists in this paper.'}), 409
    if qtype == 'MCQ' and len(answer) != 1:
        return jsonify({'message': 'An MCQ has exactly one correct option. Switch the type to MSQ for more.'}), 400

    try:
        new_id = paper_op.insert_question({
            'paper_code': data['paper_code'],
            'question_id': question_id,
            'question_type': qtype,
            'subject': data['subject'],
            'topic': data.get('topic') or None,
            'question_text': '',
            'correct_option': answer,
            'answer_text': data.get('answer_text') or None,
            'year': data.get('year') or None,
            'paper_set': data.get('paper_set') or None,
            'explanation_link': data.get('explanation_link') or None,
            'question_image': data['question_image'],
            'option_regions': _json.dumps(clean),
        })
        logger.info(f"PRACTICE IMG Q: [{admin_email}] saved image question qid={question_id} "
                    f"(db id {new_id}) in {data['paper_code']}/{data['subject']}")
        return jsonify({'message': f'Image question #{question_id} added.', 'id': new_id,
                        'question_id': question_id}), 201
    except Exception:
        logger.exception(f"PRACTICE IMG Q FAILED: [{admin_email}]")
        return jsonify({'message': 'Could not save the image question.'}), 500


@admin_bp.route('/api/practice-questions/update-image/<int:db_id>', methods=['PUT'])
@admin_login_required
def practice_update_image_question(db_id):
    import json as _json
    admin_email = session.get('admin_email', 'UnknownAdmin')
    data = request.get_json(silent=True) or {}
    existing = paper_op.get_question_by_id(db_id)
    if not existing:
        return jsonify({'message': 'Question not found.'}), 404
    required = ['question_image', 'option_regions', 'correct_option']
    missing = [f for f in required if not data.get(f)]
    if missing:
        return jsonify({'message': f"Missing: {', '.join(missing)}."}), 400
    try:
        clean = _clean_option_regions(data.get('option_regions'))
    except ValueError as e:
        return jsonify({'message': str(e)}), 400
    answer = str(data['correct_option']).strip().upper()
    for ch in answer:
        if ch.lower() not in clean:
            return jsonify({'message': f'Option {ch} is marked correct but has no area on the image.'}), 400
    qtype = str(data.get('question_type', existing.get('question_type') or 'MCQ')).upper()
    if qtype not in ('MCQ', 'MSQ'):
        return jsonify({'message': 'Image questions must be MCQ or MSQ.'}), 400

    payload = {
        'question_image': data['question_image'],
        'option_regions': _json.dumps(clean),
        'question_type': qtype,
        'correct_option': answer,
    }
    for opt_field in ('subject', 'topic', 'year', 'paper_set', 'answer_text', 'explanation_link'):
        if opt_field in data:
            payload[opt_field] = data[opt_field] or None
    try:
        paper_op.update_question(db_id, payload)
        logger.info(f"PRACTICE IMG Q EDIT: [{admin_email}] updated image question db id {db_id}")
        return jsonify({'message': 'Image question updated.'})
    except Exception:
        logger.exception(f"PRACTICE IMG Q EDIT FAILED: [{admin_email}] id={db_id}")
        return jsonify({'message': 'Could not update the image question.'}), 500


@admin_bp.route('/manage-practice-papers')
@admin_login_required
def manage_practice_papers():
    admin_email = session.get('admin_email', 'UnknownAdmin')
    try:
        from admin.manage_modules.content_db import UserOperation as ContentOperation
        content_op = ContentOperation()
        course_codes = content_op.get_distinct_course_codes()
        section_options = content_op.get_enum_values('practice_questions', 'section')
        notes_section_ids = content_op.get_notes_section_ids()
        notes_type_options = content_op.get_enum_values('notes_content', 'type')
        logger.info(f"PRACTICE PAGE: [{admin_email}] loaded with {len(course_codes)} course(s), "
                    f"{len(section_options)} format section(s), {len(notes_section_ids)} notes section(s)")
        return render_template('admin/manage_qpractice/practice_paper.html',
                               course_codes=course_codes,
                               section_options=section_options,
                               notes_section_ids=notes_section_ids,
                               notes_type_options=notes_type_options)
    except Exception as e:
        logger.exception(f"PRACTICE PAGE FAILED: [{admin_email}] {e}")
        flash("An internal server error occurred while loading the page.", "admin_error")
        return redirect(url_for('admin.admin_dashboard'))

@admin_bp.route('/api/papers/codes', methods=['GET'])
@admin_login_required
def get_paper_codes():
    try:
        codes = paper_op.get_all_paper_codes()
        return jsonify([code['code'] for code in codes])
    except Exception as e:
        logging.error(f"Error fetching paper codes: {e}", exc_info=True)
        return jsonify({'message': 'Could not load courses. Please refresh the page.'}), 500

@admin_bp.route('/api/papers/next-id', methods=['GET'])
@admin_login_required
def practice_next_question_id():
    paper_code = (request.args.get('paper_code') or '').strip()
    if not paper_code:
        return jsonify({'message': 'Paper Code is required.'}), 400
    return jsonify({'next_question_id': paper_op.next_question_id(paper_code)})


@admin_bp.route('/api/papers/subjects', methods=['GET'])
@admin_login_required
def get_subjects_for_paper():
    paper_code = request.args.get('paper_code')
    if not paper_code:
        return jsonify({'message': 'Paper Code is required.'}), 400
    try:
        subjects = paper_op.get_distinct_values('subject', paper_code)
        return jsonify([s['subject'] for s in subjects])
    except Exception as e:
        logging.error(f"Error fetching subjects for {paper_code}: {e}", exc_info=True)
        return jsonify({'message': 'Could not load subjects. Please refresh the page.'}), 500

@admin_bp.route('/api/papers/topics', methods=['GET'])
@admin_login_required
def get_topics_for_subject():
    paper_code = request.args.get('paper_code')
    subject = request.args.get('subject')
    if not paper_code or not subject:
        return jsonify({'message': 'Paper Code and Subject are required.'}), 400
    try:
        topics = paper_op.get_distinct_values('topic', paper_code, subject)
        return jsonify([t['topic'] for t in topics])
    except Exception as e:
        logging.error(f"Error fetching topics for {paper_code}/{subject}: {e}", exc_info=True)
        return jsonify({'message': 'Could not load topics. Please refresh the page.'}), 500

@admin_bp.route('/api/upload-quill-image', methods=['POST'])
@admin_login_required
def practice_upload_quill_image():
    if 'image' not in request.files:
        return jsonify({'error': 'No image file found.'}), 400
    file = request.files['image']
    if file.filename == '' or not allowed_file(file.filename):
        return jsonify({'error': 'Invalid file type.'}), 400
    try:
        filename = secure_filename(file.filename)
        timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
        unique_filename = f"{timestamp}_{filename}"
        filepath = os.path.join(UPLOAD_FOLDER, unique_filename)
        file.save(filepath)
        publish_upload(filepath)
        url_path = os.path.join('uploads', 'images', 'practice_paper', unique_filename).replace('\\', '/')
        return jsonify({'url': f'/static/{url_path}'})
    except Exception as e:
        logging.error(f"Error in practice_upload_quill_image: {e}", exc_info=True)
        return jsonify({'error': 'Could not upload the image. Please try again.'}), 500

@admin_bp.route('/api/questions/practice-csv-template', methods=['GET'])
@admin_login_required
def practice_csv_template():
    admin_email = session.get('admin_email', 'UnknownAdmin')
    logger.info(f"PRACTICE TEMPLATE: [{admin_email}] downloaded the CSV template")
    out = StringIO()
    writer = csv.writer(out)
    writer.writerow(PRACTICE_COLUMNS)
    for row in PRACTICE_TEMPLATE_ROWS:
        writer.writerow(row)
    return Response(out.getvalue(), mimetype='text/csv',
                    headers={'Content-Disposition': 'attachment; filename=practice_questions_template.csv'})


@admin_bp.route('/api/questions/practice-xlsx-template', methods=['GET'])
@admin_login_required
def practice_xlsx_template():
    admin_email = session.get('admin_email', 'UnknownAdmin')
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        logger.error("PRACTICE TEMPLATE: openpyxl is not installed")
        return jsonify({'message': 'Excel support is not installed on the server. Use the CSV template.'}), 501

    logger.info(f"PRACTICE TEMPLATE: [{admin_email}] downloaded the Excel template")
    wb = Workbook()
    ws = wb.active
    ws.title = 'Questions'
    ws.append(PRACTICE_COLUMNS)
    for row in PRACTICE_TEMPLATE_ROWS:
        ws.append(row)
    header_fill = PatternFill('solid', fgColor='833D17')
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
    for idx, width in enumerate([12, 14, 10, 12, 16, 18, 55, 18, 18, 18, 18, 14, 40, 20], 1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width

    guide = wb.create_sheet('How to fill')
    guide.append(['Column', 'Required', 'What to put'])
    for row in PRACTICE_GUIDE:
        guide.append(list(row))
    for cell in guide[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
    guide.column_dimensions['A'].width = 22
    guide.column_dimensions['B'].width = 14
    guide.column_dimensions['C'].width = 95

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='practice_questions_template.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@admin_bp.route('/api/papers/structure', methods=['GET'])
@admin_login_required
def practice_paper_structure():
    admin_email = session.get('admin_email', 'UnknownAdmin')
    paper_code = request.args.get('paper_code', '').strip()
    if not paper_code:
        return jsonify({'message': 'Paper Code is required.'}), 400
    try:
        logger.info(f"PRACTICE STRUCTURE 1: [{admin_email}] paper={paper_code!r}")
        subjects = paper_op.get_paper_structure(paper_code)
        for s in subjects:
            for k in ('questions', 'topics', 'mcq', 'msq', 'nat', 'years'):
                s[k] = int(s[k] or 0)
        totals = {
            'questions': sum(s['questions'] for s in subjects),
            'subjects': len(subjects),
            'mcq': sum(s['mcq'] for s in subjects),
            'msq': sum(s['msq'] for s in subjects),
            'nat': sum(s['nat'] for s in subjects),
        }
        logger.info(f"PRACTICE STRUCTURE 2: [{admin_email}] paper={paper_code!r} totals={totals}")
        return jsonify({'paper_code': paper_code, 'subjects': subjects, 'totals': totals})
    except Exception:
        logger.exception(f"PRACTICE STRUCTURE FAILED: [{admin_email}] paper={paper_code!r}")
        return jsonify({'message': 'Could not load the paper overview.'}), 500


@admin_bp.route('/api/questions/export-practice', methods=['GET'])
@admin_login_required
def practice_export_questions():
    admin_email = session.get('admin_email', 'UnknownAdmin')
    paper_code = request.args.get('paper_code', '').strip()
    subject = request.args.get('subject', '').strip()
    topic = request.args.get('topic', '').strip()
    fmt = (request.args.get('fmt') or 'xlsx').lower()
    if not paper_code:
        return jsonify({'message': 'Paper Code is required.'}), 400
    try:
        logger.info(f"PRACTICE EXPORT 1: [{admin_email}] paper={paper_code!r} subject={subject!r} topic={topic!r} fmt={fmt}")
        questions = paper_op.get_questions_for_export(paper_code, subject or None, topic or None)

        def row_values(q):
            return [q.get(col) if q.get(col) is not None else '' for col in PRACTICE_COLUMNS]

        safe = re.sub(r'[^A-Za-z0-9_-]', '', paper_code) or 'paper'

        if fmt == 'csv':
            out = StringIO()
            writer = csv.writer(out)
            writer.writerow(PRACTICE_COLUMNS)
            for q in questions:
                writer.writerow(row_values(q))
            logger.info(f"PRACTICE EXPORT 2: [{admin_email}] csv with {len(questions)} row(s)")
            return Response(out.getvalue(), mimetype='text/csv',
                            headers={'Content-Disposition': f'attachment; filename=practice_{safe}.csv'})

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            return jsonify({'message': 'Excel support is not installed. Use ?fmt=csv instead.'}), 501

        wb = Workbook()
        ws = wb.active
        ws.title = 'Questions'
        ws.append(PRACTICE_COLUMNS)
        for q in questions:
            ws.append(row_values(q))
        header_fill = PatternFill('solid', fgColor='833D17')
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = header_fill
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        logger.info(f"PRACTICE EXPORT 2: [{admin_email}] xlsx with {len(questions)} row(s)")
        return send_file(buf, as_attachment=True, download_name=f'practice_{safe}.xlsx',
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    except Exception:
        logger.exception(f"PRACTICE EXPORT FAILED: [{admin_email}] paper={paper_code!r}")
        return jsonify({'message': 'Could not export the questions.'}), 500


@admin_bp.route('/api/questions/upload-csv', methods=['POST'])
@admin_login_required
def practice_upload_questions_csv():
    admin_email = session.get('admin_email', 'UnknownAdmin')
    if 'csv_file' not in request.files:
        return jsonify({"message": "No file provided."}), 400
    paper_code = (request.form.get('paper_code') or '').strip()
    if not paper_code:
        return jsonify({"message": "Paper Code is required."}), 400
    csv_file = request.files['csv_file']
    if not csv_file or not csv_file.filename or not csv_file.filename.lower().endswith(('.csv', '.xlsx')):
        return jsonify({"message": "Invalid file. Please upload a .csv or .xlsx file."}), 400
    try:
        logger.info(f"PRACTICE UPLOAD 1: [{admin_email}] paper={paper_code!r} file={csv_file.filename!r}")
        reader, parse_error = _rows_from_upload(csv_file)
        if parse_error:
            logger.warning(f"PRACTICE UPLOAD BLOCKED: [{admin_email}] {parse_error}")
            return jsonify({"message": parse_error}), 400
        logger.info(f"PRACTICE UPLOAD 2: [{admin_email}] parsed {len(reader)} data row(s)")

        valid_rows, failed, errors, skipped = [], 0, [], []
        seen_ids, seen_fps = set(), {}
        by_fp, by_qid = stored_paper_index(paper_code)
        next_free = paper_op.next_question_id(paper_code)
        for i, row in enumerate(reader, 2):
            if not any(str(v or '').strip() for v in row.values()):
                continue
            try:
                question_data = {col: str(row.get(col) if row.get(col) is not None else '').strip() for col in PRACTICE_COLUMNS}
                question_data['paper_code'] = paper_code

                missing = [k for k in ('question_type', 'subject', 'correct_option', 'question_text')
                           if not question_data.get(k)]
                if missing:
                    raise ValueError(f"Missing required data: {', '.join(missing)}.")

                question_data['question_type'], question_data['correct_option'] = _validated_answer(question_data)

                if question_data['question_id'] == '':
                    while next_free in by_qid or next_free in seen_ids:
                        next_free += 1
                    qid = next_free
                    next_free += 1
                else:
                    try:
                        qid = int(float(question_data['question_id']))
                        if qid <= 0:
                            raise ValueError
                    except (TypeError, ValueError):
                        raise ValueError(f"question_id must be a positive number like 1, 2, 3 (got '{question_data['question_id']}'), or empty for the next number.")
                question_data['question_id'] = qid
                fp = paper_fingerprint(question_data)
                if fp in by_fp:
                    skipped.append({'row': i, 'error': f"already in this paper as #{by_fp[fp]['question_id']} - skipped."})
                    continue
                if fp in seen_fps:
                    skipped.append({'row': i, 'error': f'same as row {seen_fps[fp]} of this file - skipped.'})
                    continue
                if qid in by_qid:
                    raise ValueError(f"question_id '{qid}' already exists in paper {paper_code} with a different question. Use a new ID or leave it empty.")
                seen_fps[fp] = i

                if question_data.get('year'):
                    try:
                        question_data['year'] = int(float(question_data['year']))
                    except (TypeError, ValueError):
                        raise ValueError(f"year must be a number like 2024 (got '{question_data['year']}').")

                if qid in seen_ids:
                    raise ValueError(f"Duplicate question_id '{qid}' inside this file.")
                seen_ids.add(qid)

                for col in PRACTICE_COLUMNS:
                    if question_data.get(col) == '':
                        question_data[col] = None
                valid_rows.append(question_data)
            except Exception as e:
                failed += 1
                errors.append({"row": i, "error": str(e)})

        if errors:
            logger.warning(f"PRACTICE UPLOAD REFUSED: [{admin_email}] {len(errors)} problem row(s), nothing saved")
            return jsonify({"message": f"Nothing was saved - fix the {len(errors)} problem row(s) below and upload the same file again.",
                            "inserted": 0, "failed": len(errors), "errors": errors, "skipped": skipped}), 400
        if not valid_rows:
            if skipped:
                return jsonify({"message": f"Nothing added - all {len(skipped)} question(s) in this file are already in {paper_code}.",
                                "inserted": 0, "failed": 0, "errors": [], "skipped": skipped}), 200
            return jsonify({"message": "The file has no question rows."}), 400

        inserted = 0
        if valid_rows:
            try:
                inserted = paper_op.insert_questions_bulk(valid_rows)
                logger.info(f"PRACTICE UPLOAD 3: [{admin_email}] bulk-inserted {inserted} row(s) atomically")
            except Exception as exc:
                logger.exception(f"PRACTICE UPLOAD FAILED: [{admin_email}] bulk insert rolled back")
                from admin.db_errors import friendly_db_error
                exact = friendly_db_error(exc)
                return jsonify({
                    "message": (exact or "The file was read but saving failed.")
                               + " NO questions were imported - fix the issue and upload the same file again.",
                    "inserted": 0, "failed": failed + len(valid_rows), "errors": errors
                }), 500

        logger.info(f"PRACTICE UPLOAD 4: [{admin_email}] done. inserted={inserted} skipped={len(skipped)}")
        message = f"{inserted} question(s) added to {paper_code}."
        if skipped:
            message += f" {len(skipped)} already in this paper were skipped."
        return jsonify({"message": message, "inserted": inserted, "failed": 0, "errors": [], "skipped": skipped}), 200
    except Exception as e:
        logger.exception(f"PRACTICE UPLOAD FAILED: [{admin_email}] {e}")
        return jsonify({'message': 'Failed to process the file. Please check the format and try again.'}), 500

@admin_bp.route('/api/questions/add-manual', methods=['POST'])
@admin_login_required
def practice_add_manual_question():
    admin_email = session.get('admin_email', 'UnknownAdmin')
    data = request.get_json(silent=True) or {}
    logger.info(f"PRACTICE ADD 1: [{admin_email}] paper={data.get('paper_code')!r} qid={data.get('question_id')!r}")

    if not all(data.get(k) for k in ['paper_code', 'subject', 'question_type', 'correct_option']):
        logger.warning(f"PRACTICE ADD 2 BLOCKED: [{admin_email}] missing required fields")
        return jsonify({'message': 'Paper Code, Subject, Question Type and Correct Option are required.'}), 400
    if not paper_op.paper_code_exists(data['paper_code']):
        return jsonify({'message': 'Unknown paper code.'}), 400

    try:
        data['question_type'], data['correct_option'] = _validated_answer(data)
    except ValueError as e:
        logger.warning(f"PRACTICE ADD 3 BLOCKED: [{admin_email}] {e}")
        return jsonify({'message': str(e)}), 400
    if not re.sub(r'<[^>]+>|&nbsp;|\s', '', str(data.get('question_text') or '')) and '<img' not in str(data.get('question_text') or ''):
        return jsonify({'message': 'Write the question text.'}), 400

    if str(data.get('question_id') or '').strip() == '':
        data['question_id'] = paper_op.next_question_id(data['paper_code'])
    else:
        try:
            data['question_id'] = int(float(data['question_id']))
            if data['question_id'] <= 0:
                raise ValueError
        except (TypeError, ValueError):
            logger.warning(f"PRACTICE ADD 3b BLOCKED: [{admin_email}] non-numeric question_id {data.get('question_id')!r}")
            return jsonify({'message': 'Question ID must be a positive number like 1, 2, 3, or left empty for the next number.'}), 400

    if data.get('year'):
        try:
            data['year'] = int(float(data['year']))
        except (TypeError, ValueError):
            logger.warning(f"PRACTICE ADD 3c BLOCKED: [{admin_email}] non-numeric year {data.get('year')!r}")
            return jsonify({'message': 'Year must be a number like 2024, or left empty.'}), 400
    else:
        data['year'] = None
    for col in ('topic', 'paper_set', 'answer_text', 'explanation_link'):
        if data.get(col) == '':
            data[col] = None

    try:
        by_fp, by_qid = stored_paper_index(data['paper_code'])
        twin = by_fp.get(paper_fingerprint(data))
        if twin:
            logger.warning(f"PRACTICE ADD 4d BLOCKED: [{admin_email}] duplicate content of #{twin['question_id']} in {data['paper_code']!r}")
            return jsonify({'message': f"This exact question is already in this paper as #{twin['question_id']} ({twin['subject']}). Nothing was added."}), 409
        if data['question_id'] in by_qid:
            logger.warning(f"PRACTICE ADD 4 BLOCKED: [{admin_email}] duplicate qid={data['question_id']!r} in {data['paper_code']!r}")
            return jsonify({'message': f"Question ID '{data['question_id']}' already exists in this paper. Use a new ID, leave it empty for the next number, or edit the existing question."}), 409

        question_db_id = paper_op.insert_question(data)
        logger.info(f"PRACTICE ADD 5: [{admin_email}] saved question #{data['question_id']} (db_id={question_db_id})")
        return jsonify({'message': f"Question #{data['question_id']} added.", 'id': question_db_id,
                        'question_id': data['question_id'], 'next_question_id': data['question_id'] + 1}), 201
    except Exception:
        logger.exception(f"PRACTICE ADD FAILED: [{admin_email}]")
        return jsonify({'message': 'Could not save the question. Please try again.'}), 500

@admin_bp.route('/api/questions/search', methods=['GET'])
@admin_login_required
def practice_search_questions_api():
    paper_code = request.args.get('paper_code')
    subject = request.args.get('subject')
    topic = request.args.get('topic')
    if not paper_code or not subject:
        return jsonify({'message': 'Paper Code and Subject are required for search.'}), 400
    try:
        questions = paper_op.search_questions(paper_code, subject, topic)
        return jsonify({'questions': questions, 'count': len(questions)})
    except Exception as e:
        logging.error(f"Error searching questions: {e}", exc_info=True)
        return jsonify({'message': 'Search failed. Please try again.'}), 500

@admin_bp.route('/api/questions/update/<int:id>', methods=['PUT'])
@admin_login_required
def practice_update_question_api(id):
    data = request.get_json(silent=True) or {}
    if not data:
        return jsonify({'message': 'No data provided for update.'}), 400
    admin_email = session.get('admin_email', 'UnknownAdmin')

    try:
        existing = paper_op.get_question_by_id(id)
    except Exception:
        logger.exception(f"PRACTICE UPDATE FAILED: [{admin_email}] could not load db_id={id}")
        return jsonify({'message': 'Could not update the question.'}), 500

    if not existing:
        logger.warning(f"PRACTICE UPDATE BLOCKED: [{admin_email}] question db_id={id} not found")
        return jsonify({'message': 'That question no longer exists. Refresh the list and try again.'}), 404

    if not existing.get('question_image'):
        merged = dict(existing, **data)
        try:
            data['question_type'], data['correct_option'] = _validated_answer(merged)
        except ValueError as e:
            logger.warning(f"PRACTICE UPDATE BLOCKED: [{admin_email}] {e}")
            return jsonify({'message': str(e)}), 400
        by_fp, _ = stored_paper_index(existing['paper_code'])
        twin = by_fp.get(paper_fingerprint(merged))
        if twin and twin['id'] != id:
            return jsonify({'message': f"Another question in this paper already has this text and options (#{twin['question_id']})."}), 409

    if 'question_id' in data:
        try:
            data['question_id'] = int(float(data['question_id']))
            if data['question_id'] <= 0:
                raise ValueError
        except (TypeError, ValueError):
            logger.warning(f"PRACTICE UPDATE BLOCKED: [{admin_email}] bad question_id {data.get('question_id')!r}")
            return jsonify({'message': 'Question ID must be a positive number like 1, 2, 3.'}), 400

    if 'year' in data and data.get('year') not in (None, ''):
        try:
            data['year'] = int(float(data['year']))
        except (TypeError, ValueError):
            logger.warning(f"PRACTICE UPDATE BLOCKED: [{admin_email}] bad year {data.get('year')!r}")
            return jsonify({'message': 'Year must be a number like 2024, or left empty.'}), 400
    elif 'year' in data:
        data['year'] = None

    paper_code = (data.get('paper_code') or existing.get('paper_code') or '').strip()
    if 'paper_code' in data:
        if not paper_code:
            return jsonify({'message': 'Paper Code cannot be empty.'}), 400
        if not paper_op.paper_code_exists(paper_code):
            logger.warning(f"PRACTICE UPDATE BLOCKED: [{admin_email}] unknown paper_code {paper_code!r}")
            return jsonify({'message': f"Paper Code '{paper_code}' does not exist."}), 400
        data['paper_code'] = paper_code

    target_qid = data.get('question_id', existing.get('question_id'))
    try:
        if paper_op.question_id_taken_by_other(paper_code, target_qid, id):
            logger.warning(f"PRACTICE UPDATE BLOCKED: [{admin_email}] duplicate qid={target_qid!r} in {paper_code!r}")
            return jsonify({'message': f"Question ID '{target_qid}' already exists in paper {paper_code}. Use a different ID."}), 409

        logger.info(f"PRACTICE UPDATE 1: [{admin_email}] question db_id={id}")
        paper_op.update_question(id, data)
        logger.info(f"PRACTICE UPDATE 2: [{admin_email}] question db_id={id} saved")
        return jsonify({'message': 'Question updated successfully!'})
    except Exception:
        logger.exception(f"PRACTICE UPDATE FAILED: [{admin_email}] db_id={id}")
        return jsonify({'message': 'Could not update the question.'}), 500

@admin_bp.route('/api/questions/delete/<int:id>', methods=['DELETE'])
@admin_login_required
def practice_delete_question_api(id):
    if not can_manage():
        return jsonify({'success': False, 'message': 'Permission Denied.'}), 403
    admin_email = session.get('admin_email', 'UnknownAdmin')
    try:
        row = paper_op.get_question_by_id(id)
        if not row:
            return jsonify({'message': 'That question no longer exists. Refresh the list.'}), 404
        paper_op.delete_question(id)
        logger.warning(f"PRACTICE DELETE 2: [{admin_email}] question #{row['question_id']} of {row['paper_code']} deleted")
        return jsonify({'message': f"Question #{row['question_id']} deleted."})
    except Exception:
        logger.exception(f"PRACTICE DELETE FAILED: [{admin_email}] db_id={id}")
        return jsonify({'message': 'Could not delete the question.'}), 500
