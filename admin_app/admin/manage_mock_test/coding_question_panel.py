import logging
import os
import re
from datetime import datetime

from flask import jsonify, request, session
from werkzeug.utils import secure_filename

from admin import admin_bp
from admin.upload_utils import is_real_image
from admin.manage_coding.coding_db import (ASSERTION_KINDS, LANGUAGES,
                                           SECTION_TYPES, CodingOperation)
from admin.manage_mock_test.test_details_db import UserOperation
from admin.manage_practice_coding.practice_admin import (_memory_limit, question_fingerprint,
                                                         unescape_cell)
from config import Config
from extensions import admin_login_required

logger = logging.getLogger(__name__)
coding_op = CodingOperation()
user_op = UserOperation()

from admin.shared_files import publish as publish_upload
UPLOAD_FOLDER = os.path.join(Config.UPLOAD_BASE, 'coding_questions')
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
ALLOWED_IMAGE_EXT = {'png', 'jpg', 'jpeg', 'gif', 'webp'}
IMAGE_SLOTS = ('statement_image', 'input_image', 'example_image', 'schema_image_path')
DIFFICULTIES = ('Easy', 'Medium', 'Hard')


def _err(message, status=400):
    return jsonify({'success': False, 'message': message}), status


def _text(value, limit=None):
    text = value if isinstance(value, str) else ('' if value is None else str(value))
    text = text.strip()
    if limit and len(text) > limit:
        raise ValueError(f'Value is too long (limit {limit} characters).')
    return text or None


def _int(value, default, low, high, label):
    try:
        number = int(value) if value not in (None, '') else default
    except (TypeError, ValueError):
        raise ValueError(f'{label} must be a whole number.')
    if number < low or number > high:
        raise ValueError(f'{label} must be between {low} and {high}.')
    return number


def _slugify(text):
    slug = re.sub(r'[^a-z0-9]+', '-', (text or '').lower()).strip('-')
    return slug[:100] or 'coding-test'


TCS_CODE_SUFFIX = '_code'


def tcs_exam_for(description):
    return coding_op._fetch_one("SELECT * FROM tcs_exams WHERE aptitude_test_key = %s ORDER BY id DESC LIMIT 1",
                                (description.get('test_key'),))


def tcs_coding_description(description):
    return {
        'id': description.get('id'),
        'test_id': description.get('test_id'),
        'code': description.get('code'),
        'test_key': f"{description['test_key']}{TCS_CODE_SUFFIX}",
        'subject_title': f"{description.get('subject_title') or 'TCS exam'} - Coding",
        'subject_subtitle': 'Coding part of the TCS exam',
        'total_duration_minutes': 30,
    }


def ensure_tcs_exam(description):
    exam = tcs_exam_for(description)
    if exam:
        cp_test = coding_op.get_test(exam['coding_test_id'])
        if cp_test:
            return exam, cp_test
    cp_test = ensure_coding_test(tcs_coding_description(description))
    if exam:
        coding_op._execute("UPDATE tcs_exams SET coding_test_id = %s WHERE id = %s", (cp_test['id'], exam['id']))
    else:
        coding_op._execute("""
            INSERT INTO tcs_exams (section_id, code, title, subtitle, aptitude_test_key, coding_test_id, coding_label, is_published)
            VALUES (%s, %s, %s, %s, %s, %s, 'Coding', 0)
        """, (description['test_id'], description.get('code'), description.get('subject_title') or 'TCS exam',
              description.get('subject_subtitle') or None, description['test_key'], cp_test['id']))
        logger.info("CODING PANEL: created TCS exam link for description %s -> cp_tests %s", description.get('id'), cp_test['id'])
    return tcs_exam_for(description), cp_test


def _resolve_description(description_id):
    description = user_op.get_test_description_by_id(description_id)
    if not description:
        return None, _err('Test description not found.', 404)
    section = user_op.section_for_test_id(description.get('test_id'))
    if user_op.platform_for_test_id(description.get('test_id')) != 'placement':
        return None, _err('Coding questions are only allowed in the Placement Exam Portal. '
                          'Competitive courses (GATE, JEE, NEET) never carry coding tests.', 403)
    if section == 'tcs_test':
        return tcs_coding_description(description), None
    if section != 'coding_test':
        return None, _err('This description is not a Coding Test. '
                          'Set the content section to Coding Test to author coding questions.', 400)
    return description, None


def ensure_coding_test(description):
    existing = coding_op._fetch_one(
        "SELECT * FROM cp_tests WHERE test_key = %s LIMIT 1", (description['test_key'],))
    if existing:
        return existing

    slug = _slugify(f"{description.get('subject_title')}-{description['test_key'][:8]}")
    base, suffix = slug, 2
    while coding_op.slug_exists(slug):
        slug = f'{base}-{suffix}'
        suffix += 1

    test_id = coding_op._execute("""
        INSERT INTO cp_tests (slug, test_id, code, test_key, title, description,
                              total_duration_min, is_published)
        VALUES (%s, %s, %s, %s, %s, %s, %s, 1)
    """, (slug, description.get('test_id'), description.get('code'), description['test_key'],
          description.get('subject_title') or 'Coding Test',
          description.get('subject_subtitle'),
          description.get('total_duration_minutes') or 60), last_id=True)
    logger.info("CODING PANEL: created cp_tests row %s for description %s",
                test_id, description.get('id'))
    return coding_op.get_test(test_id)


@admin_bp.route('/api/coding-panel/<int:description_id>', methods=['GET'])
@admin_login_required
def coding_panel_bootstrap(description_id):
    description, error = _resolve_description(description_id)
    if error:
        return error
    try:
        cp_test = ensure_coding_test(description)
        sections = coding_op.list_sections(cp_test['id'])
    except Exception:
        logger.exception('CODING PANEL: bootstrap failed for description %s', description_id)
        return _err('Could not open the coding editor.', 500)
    return jsonify({
        'success': True,
        'test': cp_test,
        'sections': sections,
        'section_types': list(SECTION_TYPES),
        'languages': list(LANGUAGES),
        'assertion_kinds': list(ASSERTION_KINDS),
        'difficulties': list(DIFFICULTIES),
    })


@admin_bp.route('/api/coding-panel/<int:description_id>/languages', methods=['PUT'])
@admin_login_required
def coding_panel_set_languages(description_id):
    description, error = _resolve_description(description_id)
    if error:
        return error
    data = request.get_json(silent=True) or {}
    raw = data.get('languages') or []
    if not isinstance(raw, list):
        return _err('Send the languages as a list.')
    langs = []
    for item in raw:
        lang = str(item or '').strip().lower()
        if lang and lang not in langs:
            langs.append(lang)
    bad = [l for l in langs if l not in LANGUAGES]
    if bad:
        return _err(f"Unknown language: {', '.join(bad)}.")
    try:
        cp_test = ensure_coding_test(description)
        coding_op.set_test_languages(cp_test['id'], langs)
    except Exception:
        logger.exception('CODING PANEL: set languages failed for description %s', description_id)
        return _err('Could not save the languages.', 500)
    logger.info("CODING PANEL: [%s] test %s languages=%s", session.get('admin_email'),
                cp_test['id'], ','.join(langs) or 'all')
    return jsonify({'success': True, 'languages': langs})


@admin_bp.route('/api/coding-panel/<int:description_id>/sections', methods=['POST'])
@admin_login_required
def coding_panel_add_section(description_id):
    description, error = _resolve_description(description_id)
    if error:
        return error
    data = request.get_json(silent=True) or {}
    section_type = (data.get('section_type') or '').strip().lower()
    if section_type not in SECTION_TYPES:
        return _err(f"Section type must be one of: {', '.join(SECTION_TYPES)}.")
    try:
        title = _text(data.get('title'), 255)
        if not title:
            return _err('Section title is required.')
        duration = _int(data.get('duration_min'), 30, 1, 1440, 'Section duration')
    except ValueError as e:
        return _err(str(e))
    try:
        cp_test = ensure_coding_test(description)
        section_id = coding_op.create_section(cp_test['id'], title, section_type, duration)
    except Exception:
        logger.exception('CODING PANEL: add section failed')
        return _err('Could not add the section.', 500)
    logger.info("CODING PANEL: [%s] added %s section %s to description %s",
                session.get('admin_email'), section_type, section_id, description_id)
    return jsonify({'success': True, 'section_id': section_id,
                    'sections': coding_op.list_sections(cp_test['id'])})


@admin_bp.route('/api/coding-panel/sections/<int:section_id>', methods=['DELETE'])
@admin_login_required
def coding_panel_delete_section(section_id):
    section = coding_op.get_section(section_id)
    if not section:
        return _err('Section not found.', 404)
    ids = [q['id'] for q in coding_op.list_questions(section_id)]
    work = student_work(ids)
    if work['submissions'] or work['answers']:
        return _err(_work_message(work, 'a question in this section'), 409)
    try:
        coding_op.delete_section(section_id)
        note = unpublish_if_empty(section)
    except Exception:
        logger.exception('CODING PANEL: delete section %s failed', section_id)
        return _err('Could not delete the section.', 500)
    return jsonify({'success': True, 'message': 'Section deleted.' + note})


def _build_question_payload(data, section_type):
    title = _text(data.get('title'), 255)
    if not title:
        raise ValueError('Question title is required.')

    statement = _text(data.get('statement_html'))
    statement_image = _text(data.get('statement_image'), 255)
    if not statement and not statement_image:
        raise ValueError('Add the question text, or upload a question image.')

    difficulty = {'moderate': 'Medium'}.get((data.get('difficulty') or 'Medium').strip().lower(),
                                           (data.get('difficulty') or 'Medium').strip().title())
    if difficulty not in DIFFICULTIES:
        raise ValueError(f"Difficulty must be one of: {', '.join(DIFFICULTIES)}.")

    payload = {
        'title': title,
        'statement_html': statement,
        'input_spec': _text(data.get('input_spec')),
        'example_html': _text(data.get('example_html')),
        'statement_image': statement_image,
        'input_image': _text(data.get('input_image'), 255),
        'example_image': _text(data.get('example_image'), 255),
        'marks': _int(data.get('marks'), 10, 1, 1000, 'Marks'),
        'difficulty': difficulty,
        'time_limit_sec': _int(data.get('time_limit_sec'), 10, 1, 300, 'Time limit'),
        'memory_limit': _memory_limit(data.get('memory_limit')),
    }

    if section_type == 'coding':
        langs = [l.strip().lower() for l in (data.get('allowed_languages') or []) if l]
        bad = [l for l in langs if l not in LANGUAGES]
        if bad:
            raise ValueError(f"Unknown language: {', '.join(bad)}.")
        if not langs:
            raise ValueError('Select at least one allowed language.')
        test_langs = data.get('_test_languages') or []
        outside = [l for l in langs if test_langs and l not in test_langs]
        if outside:
            raise ValueError(f"This test allows only {', '.join(test_langs)}; remove {', '.join(outside)} "
                             f"or change the test's languages first.")
        payload['allowed_languages'] = ','.join(langs)
        payload['output_spec'] = _text(data.get('output_spec'))
        from admin.manage_practice_coding.practice_admin import _compare_settings
        payload['compare_mode'], payload['float_tolerance'] = _compare_settings(data)

    elif section_type == 'sql':
        sandbox = _text(data.get('sandbox_schema_sql'))
        if not sandbox:
            raise ValueError('SQL questions need the sandbox schema (CREATE TABLE + seed rows).')
        solution = _text(data.get('solution_query'))
        if not solution:
            raise ValueError('SQL questions need the solution query used to grade answers.')
        payload['sandbox_schema_sql'] = sandbox
        payload['solution_query'] = solution
        payload['schema_image_path'] = _text(data.get('schema_image_path'), 255)
        payload['ordered_comparison'] = bool(data.get('ordered_comparison'))

    elif section_type == 'web':
        payload['starter_html'] = _text(data.get('starter_html'))
        payload['starter_css'] = _text(data.get('starter_css'))
        payload['starter_js'] = _text(data.get('starter_js'))

    return payload


def test_languages_for_section(section):
    cp_test = coding_op.get_test(section['test_id']) if section else None
    raw = (cp_test or {}).get('allowed_languages') or ''
    return [l.strip() for l in raw.split(',') if l.strip()]


def section_fingerprints(section_id, exclude_id=None):
    out = {}
    for q in coding_op.list_questions(section_id):
        if exclude_id and q['id'] == exclude_id:
            continue
        out.setdefault(question_fingerprint(q['title'], q['statement_html']), q)
    return out


def student_work(question_ids):
    if not question_ids:
        return {'submissions': 0, 'answers': 0}
    marks = ','.join(['%s'] * len(question_ids))
    subs = coding_op._fetch_one(f"SELECT COUNT(*) AS n FROM cp_submissions WHERE question_id IN ({marks})", tuple(question_ids)) or {}
    ans = coding_op._fetch_one(f"SELECT COUNT(*) AS n FROM cp_answers WHERE question_id IN ({marks}) AND is_attempted = 1", tuple(question_ids)) or {}
    return {'submissions': int(subs.get('n') or 0), 'answers': int(ans.get('n') or 0)}


def _work_message(work, what):
    bits = []
    if work['submissions']:
        bits.append(f"{work['submissions']} student submission(s)")
    if work['answers']:
        bits.append(f"{work['answers']} saved answer(s)")
    return (f"Students have already worked on {what} ({', '.join(bits)}). Deleting it would erase their work, "
            f"so it cannot be removed. Edit it instead, or unpublish the exam.")


def unpublish_if_empty(section):
    cp_test = coding_op.get_test(section['test_id']) if section else None
    if not cp_test:
        return ''
    left = coding_op._fetch_one("SELECT COUNT(*) AS n FROM cp_questions q JOIN cp_sections s ON s.id = q.section_id WHERE s.test_id = %s",
                                (cp_test['id'],)) or {}
    if int(left.get('n') or 0):
        return ''
    exam = coding_op._fetch_one("SELECT id, title, is_published FROM tcs_exams WHERE coding_test_id = %s", (cp_test['id'],))
    if not exam or not exam['is_published']:
        return ''
    coding_op._execute("UPDATE tcs_exams SET is_published = 0 WHERE id = %s", (exam['id'],))
    logger.warning("CODING PANEL: exam %s unpublished - its last coding problem was deleted", exam['id'])
    return f" The exam \"{exam['title']}\" had no coding problem left, so it was unpublished; publish it again once it has one."


def _validate_children(data, section_type):
    testcases, assertions = [], []

    if section_type in ('coding', 'sql'):
        for i, case in enumerate(data.get('testcases') or [], start=1):
            expected = (case.get('expected_output') or '').strip()
            if not expected:
                raise ValueError(f'Test case {i}: expected output is required.')
            testcases.append({
                'input_text': (case.get('input_text') or '').strip(),
                'expected_output': expected,
                'is_hidden': bool(case.get('is_hidden')),
            })
        if section_type == 'coding' and not testcases:
            raise ValueError('Add at least one test case so the answer can be graded.')

    if section_type == 'web':
        for i, a in enumerate(data.get('assertions') or [], start=1):
            kind = (a.get('kind') or '').strip()
            if kind not in ASSERTION_KINDS:
                raise ValueError(f'Check {i}: choose a valid check type.')
            if not (a.get('selector') or '').strip():
                raise ValueError(f'Check {i}: a CSS selector is required.')
            if kind in ('attr_equals', 'css_property') and not (a.get('property_name') or '').strip():
                raise ValueError(f'Check {i}: "{kind}" needs an attribute/property name.')
            assertions.append(a)
        if not assertions:
            raise ValueError('Add at least one check so the page can be graded.')

    return testcases, assertions


def _save_children(question_id, section_type, data, testcases, assertions):
    if section_type in ('coding', 'sql'):
        coding_op.replace_testcases(question_id, testcases)
    if section_type == 'coding':
        coding_op.replace_templates(question_id, data.get('templates') or {})
    if section_type == 'web':
        coding_op.replace_assertions(question_id, assertions)


@admin_bp.route('/api/coding-panel/sections/<int:section_id>/questions', methods=['GET'])
@admin_login_required
def coding_panel_list_questions(section_id):
    if not coding_op.get_section(section_id):
        return _err('Section not found.', 404)
    return jsonify({'success': True, 'questions': coding_op.list_questions(section_id)})


@admin_bp.route('/api/coding-panel/starter-templates', methods=['GET'])
@admin_login_required
def coding_panel_starter_templates():
    from admin.manage_mock_test.starter_templates import STARTER_TEMPLATES
    return jsonify({'success': True, 'templates': STARTER_TEMPLATES})


@admin_bp.route('/api/coding-panel/questions/<int:question_id>', methods=['GET'])
@admin_login_required
def coding_panel_get_question(question_id):
    question = coding_op.get_question(question_id)
    if not question:
        return _err('Question not found.', 404)
    question['allowed_languages'] = [l for l in (question.get('allowed_languages') or '').split(',') if l]
    return jsonify({
        'success': True,
        'question': question,
        'testcases': coding_op.list_testcases(question_id),
        'templates': {t['language']: t['template_code'] for t in coding_op.list_templates(question_id)},
        'assertions': coding_op.list_assertions(question_id),
    })


@admin_bp.route('/api/coding-panel/sections/<int:section_id>/questions', methods=['POST'])
@admin_login_required
def coding_panel_add_question(section_id):
    section = coding_op.get_section(section_id)
    if not section:
        return _err('Section not found.', 404)
    data = request.get_json(silent=True) or {}
    section_type = section['section_type']
    data['_test_languages'] = test_languages_for_section(section)
    try:
        payload = _build_question_payload(data, section_type)
        testcases, assertions = _validate_children(data, section_type)
    except ValueError as e:
        return _err(str(e))
    already = section_fingerprints(section_id).get(question_fingerprint(payload['title'], payload['statement_html']))
    if already:
        return _err(f"This exact question is already in this section as #{already['position']} \"{already['title']}\". Nothing was added.", 409)

    payload['section_id'] = section_id
    payload['position'] = coding_op.next_position(section_id)
    try:
        question_id = coding_op.create_question(payload)
        _save_children(question_id, section_type, data, testcases, assertions)
    except Exception:
        logger.exception('CODING PANEL: add question failed for section %s', section_id)
        return _err('Could not save the question.', 500)
    logger.info("CODING PANEL: [%s] added question %s to section %s",
                session.get('admin_email'), question_id, section_id)
    return jsonify({'success': True, 'question_id': question_id})


@admin_bp.route('/api/coding-panel/questions/<int:question_id>', methods=['PUT'])
@admin_login_required
def coding_panel_update_question(question_id):
    question = coding_op.get_question(question_id)
    if not question:
        return _err('Question not found.', 404)
    data = request.get_json(silent=True) or {}
    section_type = question['section_type']
    data['_test_languages'] = test_languages_for_section(coding_op.get_section(question['section_id']))
    try:
        payload = _build_question_payload(data, section_type)
        testcases, assertions = _validate_children(data, section_type)
    except ValueError as e:
        return _err(str(e))
    already = section_fingerprints(question['section_id'], exclude_id=question_id).get(question_fingerprint(payload['title'], payload['statement_html']))
    if already:
        return _err(f"Another question in this section already has this title and statement (#{already['position']}).", 409)
    try:
        coding_op.update_question(question_id, payload)
        _save_children(question_id, section_type, data, testcases, assertions)
    except Exception:
        logger.exception('CODING PANEL: update question %s failed', question_id)
        return _err('Could not update the question.', 500)
    return jsonify({'success': True})


@admin_bp.route('/api/coding-panel/questions/<int:question_id>', methods=['DELETE'])
@admin_login_required
def coding_panel_delete_question(question_id):
    question = coding_op.get_question(question_id)
    if not question:
        return _err('Question not found.', 404)
    work = student_work([question_id])
    if work['submissions'] or work['answers']:
        return _err(_work_message(work, 'this question'), 409)
    try:
        coding_op.delete_question(question_id)
        note = unpublish_if_empty(coding_op.get_section(question['section_id']))
    except Exception:
        logger.exception('CODING PANEL: delete question %s failed', question_id)
        return _err('Could not delete the question.', 500)
    return jsonify({'success': True, 'message': 'Question deleted.' + note})


CODING_BULK_COLUMNS = ['title', 'difficulty', 'marks', 'time_limit_sec', 'memory_limit',
                       'compare_mode', 'float_tolerance',
                       'statement', 'statement_image', 'input_format', 'input_image',
                       'example', 'example_image', 'output_format', 'languages',
                       'template_python', 'template_java', 'template_c',
                       'template_cpp', 'template_csharp', 'template_javascript',
                       'tc1_input', 'tc1_output', 'tc2_input', 'tc2_output',
                       'tc3_input', 'tc3_output',
                       'hidden1_input', 'hidden1_output', 'hidden2_input', 'hidden2_output']

CODING_BULK_SAMPLE = [
    'Two Sum', 'Easy', '10', '10', '512m', 'exact', '',
    'Given n integers and a target, print the indices of the two numbers that add up to it.',
    '', 'First line: n and target. Second line: n integers.', '',
    'Input:\\n4 9\\n2 7 11 15\\n\\nOutput:\\n0 1', '',
    'Two space-separated indices.', 'python,java',
    '# read input\\nprint()', 'public class Main { }', '', '', '', '',
    '4 9\\n2 7 11 15', '0 1', '', '', '', '',
    '3 6\\n3 3 1', '0 1', '', '']

CODING_BULK_MAX_ROWS = 500


def _coding_cell(value):
    return unescape_cell(value)


def _coding_testcases_from_row(row):
    found = set()
    for key in row.keys():
        if not key:
            continue
        m = re.match(r'^(tc|hidden)(\d+)_(input|output)$', str(key).strip().lower())
        if m:
            found.add((m.group(1), int(m.group(2))))
    cases = []
    for kind, n in sorted(found, key=lambda k: (k[0] != 'tc', k[1])):
        out_val = _coding_cell(row.get(f'{kind}{n}_output'))
        if out_val:
            cases.append({'input_text': _coding_cell(row.get(f'{kind}{n}_input')),
                          'expected_output': out_val, 'is_hidden': kind == 'hidden'})
    return cases


def _coding_templates_from_row(row):
    out = {}
    for lang in LANGUAGES:
        code = row.get('template_' + lang)
        if code and str(code).strip():
            out[lang] = str(code).replace('\\n', '\n')
    return out


@admin_bp.route('/api/coding-panel/sections/<int:section_id>/bulk/template', methods=['GET'])
@admin_login_required
def coding_panel_bulk_template(section_id):
    section = coding_op.get_section(section_id)
    if not section:
        return _err('Section not found.', 404)
    if section['section_type'] != 'coding':
        return _err('Bulk import is available for coding sections only.')
    fmt = (request.args.get('fmt') or 'csv').lower()
    if fmt == 'xlsx':
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            return _err('Excel support is not installed. Use the CSV template.', 501)
        from io import BytesIO
        from flask import send_file
        wb = Workbook()
        ws = wb.active
        ws.title = 'Questions'
        ws.append(CODING_BULK_COLUMNS)
        ws.append([c.replace('\\n', '\n') for c in CODING_BULK_SAMPLE])
        fill = PatternFill('solid', fgColor='833D17')
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = fill
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, as_attachment=True,
                         download_name='coding_questions_template.xlsx',
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    import csv
    from io import StringIO
    from flask import Response
    out = StringIO()
    writer = csv.writer(out)
    writer.writerow(CODING_BULK_COLUMNS)
    writer.writerow(CODING_BULK_SAMPLE)
    return Response(out.getvalue(), mimetype='text/csv',
                    headers={'Content-Disposition':
                             'attachment; filename=coding_questions_template.csv'})


def _coding_bulk_insert_transactional(section_id, prepared):
    conn = coding_op.connection()
    conn.autocommit = False
    cur = conn.cursor()
    try:
        cur.execute("SELECT COALESCE(MAX(position),0) FROM cp_questions WHERE section_id=%s",
                    (section_id,))
        pos = (cur.fetchone() or [0])[0] or 0
        for payload, data, cases in prepared:
            pos += 1
            cur.execute("""
                INSERT INTO cp_questions
                  (section_id, position, title, statement_html, input_spec, example_html,
                   statement_image, input_image, example_image,
                   marks, difficulty, time_limit_sec, memory_limit,
                   allowed_languages, output_spec, compare_mode, float_tolerance,
                   ordered_comparison)
                VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,0)
            """, (section_id, pos, payload['title'], payload.get('statement_html'),
                  payload.get('input_spec'), payload.get('example_html'),
                  payload.get('statement_image'), payload.get('input_image'),
                  payload.get('example_image'), payload.get('marks', 10),
                  payload.get('difficulty', 'Medium'), payload.get('time_limit_sec', 10),
                  payload.get('memory_limit', '512m'), payload.get('allowed_languages'),
                  payload.get('output_spec'), payload.get('compare_mode', 'exact'),
                  payload.get('float_tolerance')))
            qid = cur.lastrowid
            if cases:
                cur.executemany(
                    "INSERT INTO cp_testcases (question_id, input_text, expected_output, "
                    "is_hidden, position) VALUES (%s,%s,%s,%s,%s)",
                    [(qid, c['input_text'], c['expected_output'],
                      1 if c['is_hidden'] else 0, i + 1) for i, c in enumerate(cases)])
            tpl_rows = [(qid, lang, code) for lang, code in (data.get('templates') or {}).items()
                        if lang in LANGUAGES and (code or '').strip()]
            if tpl_rows:
                cur.executemany(
                    "INSERT INTO cp_question_templates (question_id, language, template_code) "
                    "VALUES (%s,%s,%s)", tpl_rows)
        conn.commit()
        return len(prepared)
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/coding-panel/sections/<int:section_id>/bulk/upload', methods=['POST'])
@admin_login_required
def coding_panel_bulk_upload(section_id):
    section = coding_op.get_section(section_id)
    if not section:
        return _err('Section not found.', 404)
    if section['section_type'] != 'coding':
        return _err('Bulk import is available for coding sections only.')

    upload = request.files.get('file')
    if not upload or not upload.filename:
        return _err('Choose a CSV or Excel file to upload.')

    from admin.manage_mock_test.test_details_insert import _rows_from_upload
    rows, read_error = _rows_from_upload(upload)
    if read_error:
        return _err(read_error)
    if not rows:
        return _err('The file has no data rows.')
    if len(rows) > CODING_BULK_MAX_ROWS:
        return _err(f'Maximum {CODING_BULK_MAX_ROWS} questions per upload.')

    prepared, errors, skipped, seen = [], [], [], {}
    stored = section_fingerprints(section_id)
    test_langs = test_languages_for_section(section)
    for i, row in enumerate(rows, start=2):
        if not any(str(v or '').strip() for v in row.values()):
            continue
        try:
            languages = [l.strip().lower() for l in
                         (row.get('languages') or 'python').split(',') if l.strip()]
            data = {
                '_test_languages': test_langs,
                'title': row.get('title'), 'difficulty': row.get('difficulty') or 'Medium',
                'marks': row.get('marks'), 'time_limit_sec': row.get('time_limit_sec'),
                'memory_limit': row.get('memory_limit'),
                'statement_html': _coding_cell(row.get('statement')),
                'statement_image': _coding_cell(row.get('statement_image')),
                'input_spec': _coding_cell(row.get('input_format')),
                'input_image': _coding_cell(row.get('input_image')),
                'example_html': _coding_cell(row.get('example')),
                'example_image': _coding_cell(row.get('example_image')),
                'output_spec': _coding_cell(row.get('output_format')),
                'allowed_languages': languages,
                'compare_mode': row.get('compare_mode') or 'exact',
                'float_tolerance': row.get('float_tolerance'),
                'templates': _coding_templates_from_row(row),
                'testcases': _coding_testcases_from_row(row),
            }
            payload = _build_question_payload(data, 'coding')
            cases, _assertions = _validate_children(data, 'coding')
            fp = question_fingerprint(payload['title'], payload['statement_html'])
            if fp in stored:
                skipped.append(f"Row {i}: already in this section as #{stored[fp]['position']} - skipped.")
                continue
            if fp in seen:
                skipped.append(f'Row {i}: same as row {seen[fp]} of this file - skipped.')
                continue
            seen[fp] = i
            prepared.append((payload, data, cases))
        except ValueError as e:
            errors.append(f'Row {i}: {e}')

    if errors:
        return jsonify({'success': False,
                        'message': f'{len(errors)} row(s) need fixing. Nothing was imported.',
                        'errors': errors[:50]}), 400
    if not prepared:
        if skipped:
            return jsonify({'success': True, 'imported': 0, 'skipped': skipped,
                            'message': f'Nothing added - all {len(skipped)} question(s) in this file are already in this section.'})
        return _err('The file has no fillable rows.')

    try:
        imported = _coding_bulk_insert_transactional(section_id, prepared)
    except Exception:
        logger.exception('CODING PANEL: bulk import failed for section %s', section_id)
        return _err('A database error rolled the whole import back. Nothing was saved.', 500)

    logger.info("CODING PANEL: [%s] bulk-imported %d question(s) into section %s, %d skipped",
                session.get('admin_email'), imported, section_id, len(skipped))
    return jsonify({'success': True, 'imported': imported, 'skipped': skipped,
                    'message': f'Imported {imported} question(s).' + (f' {len(skipped)} already in this section were skipped.' if skipped else '')})


@admin_bp.route('/api/coding-panel/image', methods=['POST'])
@admin_login_required
def coding_panel_upload_image():
    slot = (request.form.get('slot') or 'statement_image').strip()
    if slot not in IMAGE_SLOTS:
        return _err('Unknown image slot.')
    file = request.files.get('image')
    if not file or not file.filename:
        return _err('No image was selected.')
    ext = file.filename.rsplit('.', 1)[-1].lower() if '.' in file.filename else ''
    if ext not in ALLOWED_IMAGE_EXT:
        return _err('Only PNG, JPG, GIF or WEBP images are allowed.')
    if not is_real_image(file):
        return _err('That file is not a valid image.')

    name = f"{datetime.now().strftime('%Y%m%d%H%M%S%f')}_{slot}_{secure_filename(file.filename)}"
    try:
        file.save(os.path.join(UPLOAD_FOLDER, name))
        publish_upload(os.path.join(UPLOAD_FOLDER, name))
    except OSError:
        logger.exception('CODING PANEL: image save failed')
        return _err('Could not save the image.', 500)
    url = '/static/' + os.path.join('uploads', 'coding_questions', name).replace('\\', '/')
    return jsonify({'success': True, 'slot': slot, 'url': url})


def sync_coding_test(description):
    test_key = (description or {}).get('test_key')
    if not test_key:
        return 0
    return coding_op._execute("""
        UPDATE cp_tests
           SET title = %s, description = %s, total_duration_min = %s
         WHERE test_key = %s
    """, (description.get('subject_title') or 'Coding Test',
          description.get('subject_subtitle'),
          description.get('total_duration_minutes') or 60,
          test_key))


def delete_coding_test_for_description(description):
    test_key = (description or {}).get('test_key')
    if not test_key:
        return 0
    exam = tcs_exam_for(description)
    if exam:
        coding_op._execute("DELETE FROM tcs_exams WHERE id = %s", (exam['id'],))
        removed = coding_op.delete_test(exam['coding_test_id'])
        logger.info("CODING PANEL: removed TCS exam %s and its coding test for description %s", exam['id'], description.get('id'))
        return removed
    removed = coding_op.delete_tests_by_key(test_key)
    if removed:
        logger.info("CODING PANEL: removed cp_tests row for deleted description %s",
                    description.get('id'))
    return removed
