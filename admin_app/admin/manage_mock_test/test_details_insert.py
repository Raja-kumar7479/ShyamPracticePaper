import csv
import hashlib
import html as _html
import re
from io import StringIO, BytesIO
import logging
import os
from datetime import datetime
from werkzeug.utils import secure_filename
from flask import render_template, request, jsonify, session, Response, send_file
from admin.manage_mock_test.test_details_db import UserOperation, COMM_ID_OFFSET
from extensions import admin_login_required, admin_owner_required, can_manage
from admin import admin_bp
import mysql.connector

logger = logging.getLogger(__name__)

user_op = UserOperation()
from config import Config
from admin.shared_files import publish as publish_upload
UPLOAD_FOLDER = os.path.join(Config.UPLOAD_BASE, 'images')
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
ALLOWED_EXTENSIONS = {'png', 'jpg', 'jpeg', 'gif'}

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def _convert_negative_marks(negative_marks_str):
    if not negative_marks_str: return 0.0
    try:
        if isinstance(negative_marks_str, (int, float)): return float(negative_marks_str)
        negative_marks_str = str(negative_marks_str).strip()
        if '/' in negative_marks_str:
            num, den = map(float, negative_marks_str.split('/'))
            if den == 0: raise ValueError("Denominator cannot be zero.")
            return num / den
        return float(negative_marks_str)
    except (ValueError, TypeError):
        raise ValueError(f'Invalid format for Negative Marks: "{negative_marks_str}". Use a number or a fraction like "1/4".')

def _norm(value):
    text = _html.unescape(re.sub(r'<[^>]+>', ' ', str(value or '')))
    return re.sub(r'\s+', ' ', text.replace('\xa0', ' ')).strip().lower()


def question_fingerprint(qtype, text, a='', b='', c='', d=''):
    body = _norm(text)
    if not body:
        return None
    raw = '|'.join([str(qtype or '').strip().upper(), body, _norm(a), _norm(b), _norm(c), _norm(d)])
    return hashlib.sha1(raw.encode('utf-8')).hexdigest()


def stored_fingerprints(description_id):
    out = {}
    for row in user_op.question_texts_for_description(description_id):
        fp = question_fingerprint(row.get('question_type'), row.get('question_text'), row.get('option_a'),
                                  row.get('option_b'), row.get('option_c'), row.get('option_d'))
        if fp and fp not in out:
            out[fp] = row
    return out


def where_stored(row):
    return f"{row.get('section_name') or row.get('section_id')} Q{row.get('question_number')}"


def _has_text(html):
    return bool(re.sub(r'<[^>]+>|&nbsp;|\s', '', html or ''))


def _check_answer(qtype, correct, options=None):
    qtype = str(qtype or '').strip().upper()
    answer = str(correct or '').strip().upper()
    if qtype not in ('MCQ', 'MSQ', 'NAT'):
        raise ValueError(f"Question type must be MCQ, MSQ or NAT (got '{qtype or 'empty'}').")
    if not answer:
        raise ValueError('The correct answer is required.')
    if qtype == 'NAT':
        try:
            float(answer)
        except ValueError:
            raise ValueError(f"A NAT answer must be a number (got '{correct}').")
        return qtype, str(correct).strip()
    letters = answer.replace(',', '').replace(' ', '')
    if qtype == 'MCQ' and (len(letters) != 1 or letters not in 'ABCD'):
        raise ValueError(f"An MCQ answer must be one letter A, B, C or D (got '{correct}').")
    if qtype == 'MSQ':
        if not letters or any(ch not in 'ABCD' for ch in letters) or len(set(letters)) != len(letters):
            raise ValueError(f"An MSQ answer must be letters from A-D without repeats, like AC (got '{correct}').")
        letters = ''.join(sorted(letters))
    if options is not None:
        for ch in letters:
            if not _has_text(options.get(ch.lower())):
                raise ValueError(f"Option {ch} is the correct answer but option {ch} is empty.")
    return qtype, letters


@admin_bp.route('/api/question-image/upload', methods=['POST'])
@admin_login_required
def upload_question_image():
    admin_email = session.get('admin_email', 'UnknownAdmin')
    file = request.files.get('image')
    if not file or not file.filename:
        return jsonify({'message': 'Choose an image file first.'}), 400
    if not allowed_file(file.filename):
        logger.warning(f"IMG UPLOAD BLOCKED: [{admin_email}] bad type {file.filename!r}")
        return jsonify({'message': 'Only PNG, JPG, JPEG or GIF images are allowed.'}), 400
    try:
        name = f"{datetime.now().strftime('%Y%m%d%H%M%S')}_{secure_filename(file.filename)}"
        file.save(os.path.join(UPLOAD_FOLDER, name))
        publish_upload(os.path.join(UPLOAD_FOLDER, name))
        url = '/static/' + os.path.join('uploads', 'images', name).replace('\\', '/')
        logger.info(f"IMG UPLOAD: [{admin_email}] saved question image {name}")
        return jsonify({'url': url})
    except Exception:
        logger.exception(f"IMG UPLOAD FAILED: [{admin_email}]")
        return jsonify({'message': 'Could not save the image. Please try again.'}), 500


@admin_bp.route('/api/test-questions/add-image', methods=['POST'])
@admin_login_required
def add_image_question_api():
    import json as _json
    admin_email = session.get('admin_email', 'UnknownAdmin')
    data = request.get_json(silent=True) or {}
    required = ['test_id', 'test_key', 'test_description_id', 'section_id', 'section_name',
                'question_image', 'option_regions', 'correct_option', 'correct_marks']
    missing = [f for f in required if not data.get(f)]
    if missing:
        logger.warning(f"IMG Q BLOCKED: [{admin_email}] missing {missing}")
        return jsonify({'message': f"Missing: {', '.join(missing)}."}), 400

    regions = data.get('option_regions')
    if not isinstance(regions, dict) or not regions:
        return jsonify({'message': 'Mark at least one option area on the image.'}), 400
    clean_regions = {}
    for key in ('a', 'b', 'c', 'd'):
        box = regions.get(key)
        if not isinstance(box, dict):
            continue
        try:
            clean_regions[key] = {k: round(float(box[k]), 2) for k in ('x', 'y', 'w', 'h')}
        except (KeyError, TypeError, ValueError):
            return jsonify({'message': f'Option {key.upper()} area is invalid. Mark it again.'}), 400
    if not clean_regions:
        return jsonify({'message': 'Mark at least one option area on the image.'}), 400

    answer = str(data['correct_option']).strip().upper()
    for ch in answer:
        if ch.lower() not in clean_regions:
            return jsonify({'message': f'Option {ch} is marked correct but has no area on the image.'}), 400

    try:
        desc = user_op.get_test_description_by_id(data['test_description_id'])
        if not desc:
            return jsonify({'message': 'Test Description not found.'}), 404
        declared = int(desc.get('total_questions') or 0)
        current = user_op.get_question_count_for_description(data['test_description_id'])
        if current >= declared:
            return jsonify({'message': f"All {declared} questions are already uploaded for this test."}), 409

        payload = dict(data)
        payload['question_text'] = ''
        payload['question_type'] = str(data.get('question_type', 'MCQ')).upper()
        payload['question_level'] = data.get('question_level') or 'Medium'
        payload['negative_marks'] = _convert_negative_marks(data.get('negative_marks'))
        payload['correct_marks'] = float(data['correct_marks'])
        payload['correct_option'] = answer
        payload['option_regions'] = _json.dumps(clean_regions)
        payload['question_number'] = user_op.get_max_question_number_in_section(
            data['test_description_id'], data['section_id']) + 1
        user_op.insert_test_question(payload)
        logger.info(f"IMG Q: [{admin_email}] saved image question {current + 1}/{declared} "
                    f"with {len(clean_regions)} option area(s)")
        return jsonify({'message': f'Image question added ({current + 1}/{declared}).'}), 201
    except ValueError as e:
        return jsonify({'message': str(e)}), 400
    except Exception as exc:
        logger.exception(f"IMG Q FAILED: [{admin_email}]")
        from admin.db_errors import friendly_db_error
        return jsonify({'message': friendly_db_error(exc) or 'Could not save the image question.'}), 500


@admin_bp.route('/api/test-questions/update-image/<int:qid>', methods=['PUT'])
@admin_login_required
def update_image_question_api(qid):
    import json as _json
    admin_email = session.get('admin_email', 'UnknownAdmin')
    data = request.get_json(silent=True) or {}
    if not user_op.question_exists(qid):
        logger.warning(f"IMG Q EDIT BLOCKED: [{admin_email}] question {qid} not found")
        return jsonify({'message': 'Question not found.'}), 404
    required = ['question_image', 'option_regions', 'correct_option', 'correct_marks']
    missing = [f for f in required if not data.get(f)]
    if missing:
        return jsonify({'message': f"Missing: {', '.join(missing)}."}), 400

    regions = data.get('option_regions')
    if not isinstance(regions, dict) or not regions:
        return jsonify({'message': 'Mark at least one option area on the image.'}), 400
    clean_regions = {}
    for key in ('a', 'b', 'c', 'd'):
        box = regions.get(key)
        if not isinstance(box, dict):
            continue
        try:
            clean_regions[key] = {k: round(float(box[k]), 2) for k in ('x', 'y', 'w', 'h')}
        except (KeyError, TypeError, ValueError):
            return jsonify({'message': f'Option {key.upper()} area is invalid. Mark it again.'}), 400
    if not clean_regions:
        return jsonify({'message': 'Mark at least one option area on the image.'}), 400

    answer = str(data['correct_option']).strip().upper()
    for ch in answer:
        if ch.lower() not in clean_regions:
            return jsonify({'message': f'Option {ch} is marked correct but has no area on the image.'}), 400

    try:
        payload = {
            'question_image': data['question_image'],
            'option_regions': _json.dumps(clean_regions),
            'question_type': str(data.get('question_type', 'MCQ')).upper(),
            'question_level': data.get('question_level') or 'Medium',
            'negative_marks': _convert_negative_marks(data.get('negative_marks')),
            'correct_marks': float(data['correct_marks']),
            'correct_option': answer,
        }
        user_op.update_image_question(qid, payload)
        logger.info(f"IMG Q EDIT: [{admin_email}] updated image question {qid} "
                    f"with {len(clean_regions)} option area(s)")
        return jsonify({'message': 'Image question updated.'})
    except ValueError as e:
        return jsonify({'message': str(e)}), 400
    except Exception as exc:
        logger.exception(f"IMG Q EDIT FAILED: [{admin_email}] id={qid}")
        from admin.db_errors import friendly_db_error
        return jsonify({'message': friendly_db_error(exc) or 'Could not update the image question.'}), 500


@admin_bp.route('/api/test-descriptions/<int:description_id>/flagged', methods=['GET'])
@admin_login_required
def get_flagged_questions_api(description_id):
    admin_email = session.get('admin_email', 'UnknownAdmin')
    try:
        rows = user_op.get_flagged_questions_for_description(description_id)
        logger.info(f"FLAGS: [{admin_email}] {len(rows)} flagged question(s) for description {description_id}")
        return jsonify(rows)
    except Exception:
        logger.exception(f"FLAGS FAILED: [{admin_email}] description={description_id}")
        return jsonify({'message': 'Could not load flagged questions.'}), 500


@admin_bp.route('/manage-tests', methods=['GET'])
@admin_login_required
def manage_tests():
    admin_email = session.get('admin_email', 'UnknownAdmin')
    logger.info(f"[{admin_email}] Accessed /manage-tests")
    return render_template('admin/manage_mock_test/test_details_insert.html')


@admin_bp.route('/manage-tests/image-batch', methods=['GET'])
@admin_login_required
def image_batch_page():
    admin_email = session.get('admin_email', 'UnknownAdmin')
    logger.info(f"[{admin_email}] Accessed /manage-tests/image-batch")
    return render_template('admin/manage_mock_test/image_batch.html')

@admin_bp.route('/api/content-for-tests', methods=['GET'])
@admin_login_required
def get_content_for_tests_api():
    admin_email = session.get('admin_email', 'UnknownAdmin')
    try:
        logger.info(f"[{admin_email}] Requested test content")
        content = user_op.get_content_for_tests()
        return jsonify(content)
    except Exception as e:
        logger.exception(f"[{admin_email}] Error in get_content_for_tests_api")
        return jsonify({'message': 'Could not load the test list. Please refresh the page.'}), 500

@admin_bp.route('/api/test-descriptions/<test_id>', methods=['GET'])
@admin_login_required
def get_test_descriptions_api(test_id):
    admin_email = session.get('admin_email', 'UnknownAdmin')
    try:
        logger.info(f"[{admin_email}] Fetching test descriptions for test_id={test_id}")
        descriptions = user_op.get_test_descriptions_by_test_id(test_id)
        if descriptions and descriptions[0].get('section') == 'tcs_test':
            for row in descriptions:
                row['tcs'] = user_op.tcs_exam_summary(row['test_key'])
        return jsonify(descriptions)
    except Exception as e:
        logger.exception(f"[{admin_email}] Failed fetching test descriptions for test_id={test_id}")
        return jsonify({'message': 'Could not load test descriptions. Please refresh the page.'}), 500

KEY_PREFIX = {'free': 'free_', 'paid': 'prem_'}
_KEY_RE = {
    'free': re.compile(r'^free_[0-9a-f]{20}$'),
    'paid': re.compile(r'^prem_[0-9a-f]{20}$'),
}


def _key_matches_access(test_key, access_label):
    pattern = _KEY_RE.get(access_label)
    if pattern is None:
        return not (test_key.startswith('free_') or test_key.startswith('prem_'))
    return bool(pattern.match(test_key or ''))


def _new_test_key(access_label):
    import secrets as _secrets
    prefix = KEY_PREFIX.get(access_label, '')
    for _ in range(8):
        candidate = prefix + _secrets.token_hex(10) if prefix else _secrets.token_hex(12)
        if not user_op.test_key_exists(candidate):
            return candidate
    return None


def _resolve_access(test_id, requested):
    if user_op.content_label_for_test_id(test_id) == 'Paid':
        return 'inherit', None
    if requested not in ('free', 'paid'):
        return None, 'Choose Free or Premium for Test Access first.'
    if requested == 'paid':
        from admin.manage_premium.premium_catalog import course_has_premium, PREMIUM_NOT_SET
        if not course_has_premium(user_op.get_code_for_test_id(test_id)):
            return None, PREMIUM_NOT_SET
    return requested, None


NO_CALCULATOR_SECTIONS = ('coding_test', 'communication_test')


def _limit_reached(test_id):
    usage = user_op.description_usage(test_id)
    if usage['limit'] is not None and usage['count'] >= usage['limit']:
        return (f"Limit reached: this content allows {usage['limit']} test description(s) and already "
                f"has {usage['count']}. Not allowed to add more - raise the limit in Course Content first.")
    return None


@admin_bp.route('/api/test-descriptions/add', methods=['POST'])
@admin_login_required
def add_test_description_api():
    admin_email = session.get('admin_email', 'UnknownAdmin')
    logger.info(f"ADD DESC 1: [{admin_email}] request received")

    data = request.get_json(silent=True) or {}
    required = ['test_id', 'test_key', 'subject_title', 'total_questions', 'total_marks', 'total_duration_minutes']
    missing = [f for f in required if not data.get(f)]
    if missing:
        logger.warning(f"ADD DESC 2 BLOCKED: [{admin_email}] missing fields {missing}")
        return jsonify({'message': f"Missing required fields: {', '.join(missing)}."}), 400

    try:
        total_questions = int(data['total_questions'])
        total_marks = float(data['total_marks'])
        total_duration = int(data['total_duration_minutes'])
    except (TypeError, ValueError):
        logger.warning(f"ADD DESC 3 BLOCKED: [{admin_email}] non-numeric values in {data}")
        return jsonify({'message': 'Questions, marks and duration must be numbers.'}), 400

    if total_questions <= 0 or total_marks <= 0 or total_duration <= 0:
        logger.warning(f"ADD DESC 4 BLOCKED: [{admin_email}] non-positive values in {data}")
        return jsonify({'message': 'Questions, marks and duration must be greater than zero.'}), 400

    try:
        code = user_op.get_code_for_test_id(data['test_id'])
        if not code:
            logger.warning(f"ADD DESC 5 BLOCKED: [{admin_email}] unknown test_id={data['test_id']!r}")
            return jsonify({'message': 'This Test ID does not belong to any course. Please reselect the test.'}), 400

        limit_msg = _limit_reached(data['test_id'])
        if limit_msg:
            logger.warning(f"ADD DESC 5L BLOCKED: [{admin_email}] {limit_msg}")
            return jsonify({'message': limit_msg, 'limit_reached': True}), 409

        if user_op.test_key_exists(data['test_key']):
            logger.warning(f"ADD DESC 6 BLOCKED: [{admin_email}] duplicate test_key={data['test_key']!r}")
            return jsonify({'message': 'That Test Key is already used. Generate a new key.'}), 409

        try:
            show_calculator = 1 if int(data.get('show_calculator', 1)) else 0
        except (TypeError, ValueError):
            show_calculator = 1
        if user_op.section_for_test_id(data['test_id']) in NO_CALCULATOR_SECTIONS:
            show_calculator = 0

        access_label, access_error = _resolve_access(data['test_id'], data.get('access_label'))
        if access_error:
            logger.warning(f"ADD DESC 6a BLOCKED: [{admin_email}] {access_error}")
            return jsonify({'message': access_error}), 400
        if not _key_matches_access(str(data['test_key']), access_label):
            logger.warning(f"ADD DESC 6b BLOCKED: [{admin_email}] key {data['test_key']!r} does not match access {access_label!r}")
            return jsonify({'message': 'Generate the Test Key again after choosing Free or Premium - '
                                       'the key is built from that choice.'}), 400

        payload = {
            'show_calculator': show_calculator,
            'access_label': access_label,
            'test_id': data['test_id'],
            'code': code,
            'test_key': data['test_key'],
            'subject_title': data['subject_title'],
            'subject_subtitle': data.get('subject_subtitle') or '',
            'total_questions': total_questions,
            'total_marks': total_marks,
            'total_duration_minutes': total_duration,
        }
        logger.info(f"ADD DESC 7: [{admin_email}] inserting {payload['test_key']} under code={code}")
        user_op.insert_test_description(payload)

        block_section = user_op.section_for_test_id(payload['test_id'])
        if block_section == 'coding_test':
            from admin.manage_mock_test.coding_question_panel import ensure_coding_test
            try:
                ensure_coding_test(payload)
            except Exception:
                logger.exception("ADD DESC 7b: [%s] could not create the coding test for %s",
                                 admin_email, payload['test_key'])
        if block_section == 'tcs_test':
            from admin.manage_mock_test.coding_question_panel import ensure_tcs_exam
            try:
                ensure_tcs_exam(payload)
            except Exception:
                logger.exception("ADD DESC 7c: [%s] could not create the TCS exam link for %s",
                                 admin_email, payload['test_key'])
            logger.info(f"ADD DESC 8: [{admin_email}] TCS exam added for test_id={data['test_id']}")
            return jsonify({'message': 'TCS exam added: upload its aptitude questions with the list button and its '
                                       'coding problems with the code button, then publish it.'}), 201

        logger.info(f"ADD DESC 8: [{admin_email}] test description added for test_id={data['test_id']}")
        return jsonify({'message': 'Test Description added successfully!'}), 201

    except ValueError as e:
        logger.warning(f"ADD DESC FAILED (validation): [{admin_email}] {e}")
        return jsonify({'message': str(e)}), 400
    except Exception:
        logger.exception(f"ADD DESC FAILED: [{admin_email}] could not add test description")
        return jsonify({'message': 'Could not save the test description. Please try again.'}), 500

@admin_bp.route('/api/test-descriptions/update/<int:id>', methods=['PUT'])
@admin_login_required
def update_test_description_api(id):
    admin_email = session.get('admin_email', 'UnknownAdmin')
    logger.info(f"UPDATE DESC 1: [{admin_email}] request for id={id}")

    data = request.get_json(silent=True) or {}
    if not data:
        logger.warning(f"UPDATE DESC 2 BLOCKED: [{admin_email}] no data for id={id}")
        return jsonify({'message': 'No data provided.'}), 400

    if 'show_calculator' in data:
        try:
            data['show_calculator'] = 1 if int(data['show_calculator']) else 0
        except (TypeError, ValueError):
            data['show_calculator'] = 1
        owner = user_op.get_test_description_by_id(id)
        if owner and user_op.section_for_test_id(owner.get('test_id')) in NO_CALCULATOR_SECTIONS:
            data['show_calculator'] = 0

    numeric_rules = {'total_questions': int, 'total_marks': float, 'total_duration_minutes': int}
    for field, caster in numeric_rules.items():
        if field in data:
            try:
                data[field] = caster(data[field])
            except (TypeError, ValueError):
                logger.warning(f"UPDATE DESC 2b BLOCKED: [{admin_email}] non-numeric {field}={data[field]!r}")
                return jsonify({'message': f"{field.replace('_', ' ').title()} must be a number."}), 400
            if data[field] <= 0:
                logger.warning(f"UPDATE DESC 2c BLOCKED: [{admin_email}] non-positive {field}={data[field]}")
                return jsonify({'message': f"{field.replace('_', ' ').title()} must be greater than zero."}), 400

    if 'total_questions' in data:
        uploaded = user_op.get_question_count_for_description(id)
        if data['total_questions'] < uploaded:
            logger.warning(f"UPDATE DESC 2d BLOCKED: [{admin_email}] id={id} tried limit {data['total_questions']} < uploaded {uploaded}")
            return jsonify({
                'message': f"This test already has {uploaded} uploaded question(s). "
                           f"Total Questions cannot be less than that - delete questions first."
            }), 400

    new_key = None
    if 'access_label' in data:
        current = user_op.get_test_description_by_id(id)
        if not current:
            return jsonify({'message': 'Test description not found.'}), 404
        access_label, access_error = _resolve_access(current.get('test_id'), data['access_label'])
        if access_error:
            return jsonify({'message': access_error}), 400
        data['access_label'] = access_label
        current_access = current.get('access_label') if current.get('access_label') in ('free', 'paid') else 'free'
        if access_label != 'inherit' and access_label != current_access:
            activity = user_op.student_activity_for_description(id)
            if activity:
                detail = ', '.join(f"{n} {k}" for k, n in activity.items())
                logger.warning(f"UPDATE DESC 2e BLOCKED: [{admin_email}] id={id} access change with student activity ({detail})")
                return jsonify({'message': f"Students have already worked on this test ({detail}), so its access "
                                           f"cannot change. Clone it and set the copy to "
                                           f"{'Premium' if access_label == 'paid' else 'Free'} instead."}), 409
            new_key = _new_test_key(access_label)
            if not new_key:
                return jsonify({'message': 'Could not generate a new test key. Try again.'}), 500

    try:
        for field, value in data.items():
            logger.info(f"UPDATE DESC 3: [{admin_email}] id={id} field={field!r}")
            user_op.update_test_description_field(id, field, value)

        if new_key:
            old_key = user_op.rekey_test_description(id, new_key)
            logger.info(f"UPDATE DESC 3k: [{admin_email}] id={id} re-keyed {old_key} -> {new_key}")

        if {'subject_title', 'subject_subtitle', 'total_duration_minutes'} & set(data):
            updated = user_op.get_test_description_by_id(id)
            if updated and user_op.section_for_test_id(updated.get('test_id')) == 'coding_test':
                from admin.manage_mock_test.coding_question_panel import sync_coding_test
                try:
                    sync_coding_test(updated)
                except Exception:
                    logger.exception("UPDATE DESC 4b: [%s] could not sync the coding test for id=%s",
                                     admin_email, id)
        if new_key or ({'subject_title', 'subject_subtitle'} & set(data)):
            updated = user_op.get_test_description_by_id(id)
            if updated and user_op.section_for_test_id(updated.get('test_id')) == 'tcs_test':
                try:
                    user_op.sync_tcs_exam(old_key if new_key else updated['test_key'], updated)
                except Exception:
                    logger.exception("UPDATE DESC 4c: [%s] could not sync the TCS exam for id=%s", admin_email, id)

        logger.info(f"UPDATE DESC 4: [{admin_email}] updated id={id} fields={list(data)}")
        if new_key:
            return jsonify({'message': f'Test Description updated. Its access changed, so it has a new key: {new_key}',
                            'new_test_key': new_key})
        return jsonify({'message': 'Test Description updated successfully!'})
    except ValueError as e:
        logger.warning(f"UPDATE DESC BLOCKED: [{admin_email}] id={id}: {e}")
        return jsonify({'message': 'That field cannot be updated.'}), 400
    except Exception:
        logger.exception(f"UPDATE DESC FAILED: [{admin_email}] id={id}")
        return jsonify({'message': 'Could not update the test description.'}), 500

@admin_bp.route('/api/test-descriptions/delete/<int:id>', methods=['DELETE'])
@admin_login_required
def delete_test_description_api(id):
    if not can_manage():
        return jsonify({'success': False, 'message': 'Permission Denied.'}), 403
    admin_email = session.get('admin_email', 'UnknownAdmin')
    logger.info(f"DELETE DESC 1: [{admin_email}] request for id={id}")
    try:
        existing = user_op.get_test_description_by_id(id)

        activity = user_op.student_activity_for_description(id)
        if existing and user_op.section_for_test_id(existing.get('test_id')) == 'tcs_test':
            tcs = user_op.tcs_exam_summary(existing['test_key'])
            if tcs and tcs.get('attempts'):
                activity['exam attempts'] = tcs['attempts']
        if activity:
            summary = ', '.join(f"{n} {label}" for label, n in activity.items())
            logger.warning(f"DELETE DESC BLOCKED: [{admin_email}] id={id} has student work ({summary})")
            return jsonify({'message': "Students have already used this test (" + summary + "). "
                                       "Deleting it would erase their attempts and results, so it cannot be "
                                       "removed. Edit its questions, or hide the content block it sits under."}), 409

        user_op.delete_test_description(id)
        if existing:
            from admin.manage_mock_test.coding_question_panel import delete_coding_test_for_description
            try:
                delete_coding_test_for_description(existing)
            except Exception:
                logger.exception("DELETE DESC 1b: [%s] could not remove the coding test for id=%s",
                                 admin_email, id)
        logger.warning(f"DELETE DESC 2: [{admin_email}] deleted description id={id} and its questions")
        return jsonify({'message': 'Test Description and all its questions deleted successfully!'})
    except Exception:
        logger.exception(f"DELETE DESC FAILED: [{admin_email}] id={id}")
        return jsonify({'message': 'Could not delete the test description.'}), 500

@admin_bp.route('/api/test-structure/<int:description_id>', methods=['GET'])
@admin_login_required
def get_test_structure_api(description_id):
    admin_email = session.get('admin_email', 'UnknownAdmin')
    try:
        logger.info(f"STRUCTURE 1: [{admin_email}] request for description_id={description_id}")

        desc = user_op.get_test_description_by_id(description_id)
        if not desc:
            logger.warning(f"STRUCTURE 2 BLOCKED: [{admin_email}] description {description_id} not found")
            return jsonify({'message': 'Test Description not found.'}), 404

        sections = user_op.get_test_structure(description_id)
        for s in sections:
            s['questions'] = int(s['questions'] or 0)
            s['mcq'] = int(s['mcq'] or 0)
            s['msq'] = int(s['msq'] or 0)
            s['nat'] = int(s['nat'] or 0)
            s['marks'] = float(s['marks'] or 0)

        actual_questions = sum(s['questions'] for s in sections)
        actual_marks = round(sum(s['marks'] for s in sections), 2)
        declared_questions = int(desc.get('total_questions') or 0)
        declared_marks = float(desc.get('total_marks') or 0)

        if actual_questions > declared_questions:
            q_status = 'over'
        elif actual_questions == declared_questions:
            q_status = 'complete'
        else:
            q_status = 'under'

        logger.info(
            f"STRUCTURE 3: [{admin_email}] id={description_id} "
            f"declared={declared_questions}q/{declared_marks}m "
            f"actual={actual_questions}q/{actual_marks}m status={q_status}"
        )

        return jsonify({
            'declared': {
                'test_key': desc.get('test_key'),
                'subject_title': desc.get('subject_title'),
                'subject_subtitle': desc.get('subject_subtitle'),
                'total_questions': declared_questions,
                'total_marks': declared_marks,
                'total_duration_minutes': int(desc.get('total_duration_minutes') or 0),
            },
            'actual': {
                'questions': actual_questions,
                'marks': actual_marks,
                'mcq': sum(s['mcq'] for s in sections),
                'msq': sum(s['msq'] for s in sections),
                'nat': sum(s['nat'] for s in sections),
            },
            'sections': sections,
            'status': {
                'questions': q_status,
                'remaining': max(0, declared_questions - actual_questions),
                'over_by': max(0, actual_questions - declared_questions),
                'marks_match': abs(actual_marks - declared_marks) < 0.01,
                'marks_diff': round(actual_marks - declared_marks, 2),
            }
        })
    except Exception:
        logger.exception(f"STRUCTURE FAILED: [{admin_email}] description_id={description_id}")
        return jsonify({'message': 'Could not load the test structure.'}), 500


CSV_TEMPLATE_COLUMNS = ['question_type', 'question_text', 'option_a', 'option_b', 'option_c', 'option_d',
                        'correct_option', 'correct_marks', 'negative_marks', 'question_level',
                        'answer_text', 'answer_link']

CSV_TEMPLATE_ROWS = [
    ['MCQ', '<p>What is the capital of India?</p>', '<p>Mumbai</p>', '<p>New Delhi</p>', '<p>Kolkata</p>', '<p>Chennai</p>',
     'B', '2', '0.5', 'Easy', '<p>New Delhi is the capital of India.</p>', ''],
    ['MSQ', '<p>Which of these are prime numbers?</p>', '<p>2</p>', '<p>4</p>', '<p>5</p>', '<p>9</p>',
     'AC', '2', '0', 'Medium', '<p>2 and 5 are prime. 4 and 9 are composite.</p>', ''],
    ['NAT', '<p>If x + 3 = 10, the value of x is ______</p>', '', '', '', '',
     '7', '1', '0', 'Easy', '<p>x = 10 - 3 = 7</p>', ''],
    ['MCQ',
     '<p>What is the output of this Python code?</p><pre class="ql-syntax">def add(a, b):\n    return a + b\n\nprint(add(2, 3))</pre>',
     '<p>23</p>', '<p>5</p>', '<p>Error</p>', '<p>None</p>',
     'B', '2', '1/4', 'Medium',
     '<p>The function returns the sum, so it prints 5.</p><pre class="ql-syntax">add(2, 3)  # 2 + 3 = 5</pre>', ''],
]


COLUMN_GUIDE = [
    ('question_type', 'Yes', 'MCQ, MSQ or NAT'),
    ('question_text', 'Yes', 'The question. Plain text or HTML. For code wrap it in <pre class="ql-syntax">...</pre>'),
    ('option_a .. option_d', 'MCQ/MSQ only', 'The four options. Leave empty for NAT'),
    ('correct_option', 'Yes', 'MCQ: one letter (B). MSQ: letters joined (AC). NAT: the numeric answer (7 or 12.5)'),
    ('correct_marks', 'Yes', 'Marks for a correct answer, e.g. 2'),
    ('negative_marks', 'No', 'Penalty as a number or fraction (0.5 or 1/4). Empty = 0'),
    ('question_level', 'No', 'Easy, Medium or Tough (default Medium)'),
    ('answer_text', 'No', 'Worked explanation shown after the test'),
    ('answer_link', 'No', 'Optional URL to a video/page explanation'),
]


@admin_bp.route('/api/test-questions/csv-template', methods=['GET'])
@admin_login_required
def download_questions_csv_template():
    admin_email = session.get('admin_email', 'UnknownAdmin')
    logger.info(f"CSV TEMPLATE: [{admin_email}] downloaded the CSV question template")
    out = StringIO()
    writer = csv.writer(out)
    writer.writerow(CSV_TEMPLATE_COLUMNS)
    for row in CSV_TEMPLATE_ROWS:
        writer.writerow(row)
    return Response(
        out.getvalue(),
        mimetype='text/csv',
        headers={'Content-Disposition': 'attachment; filename=question_upload_template.csv'}
    )


@admin_bp.route('/api/test-questions/xlsx-template', methods=['GET'])
@admin_login_required
def download_questions_xlsx_template():
    admin_email = session.get('admin_email', 'UnknownAdmin')
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        logger.error("XLSX TEMPLATE: openpyxl is not installed")
        return jsonify({'message': 'Excel support is not installed on the server. Use the CSV template.'}), 501

    logger.info(f"XLSX TEMPLATE: [{admin_email}] downloaded the Excel question template")
    wb = Workbook()
    ws = wb.active
    ws.title = 'Questions'
    ws.append(CSV_TEMPLATE_COLUMNS)
    for row in CSV_TEMPLATE_ROWS:
        ws.append(row)

    header_fill = PatternFill('solid', fgColor='833D17')
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
    widths = [14, 55, 20, 20, 20, 20, 14, 13, 14, 14, 40, 20]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width

    guide = wb.create_sheet('How to fill')
    guide.append(['Column', 'Required', 'What to put'])
    for row in COLUMN_GUIDE:
        guide.append(list(row))
    for cell in guide[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
    guide.column_dimensions['A'].width = 22
    guide.column_dimensions['B'].width = 14
    guide.column_dimensions['C'].width = 90

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='question_upload_template.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


def _rows_from_upload(upload):
    filename = (upload.filename or '').lower()

    if filename.endswith('.csv'):
        stream = StringIO(upload.stream.read().decode('utf-8-sig'))
        return list(csv.DictReader(stream)), None

    if filename.endswith('.xlsx'):
        try:
            from openpyxl import load_workbook
        except ImportError:
            return None, 'Excel support is not installed on the server. Upload a CSV instead.'
        wb = load_workbook(BytesIO(upload.stream.read()), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows_iter = ws.iter_rows(values_only=True)
        try:
            headers = [str(h).strip() if h is not None else '' for h in next(rows_iter)]
        except StopIteration:
            wb.close()
            return [], None
        rows = []
        for values in rows_iter:
            row = {}
            for idx, header in enumerate(headers):
                if not header:
                    continue
                value = values[idx] if idx < len(values) else None
                if value is None:
                    row[header] = ''
                elif isinstance(value, float) and value == int(value):
                    row[header] = str(int(value))
                else:
                    row[header] = str(value)
            rows.append(row)
        wb.close()
        return rows, None

    return None, 'Invalid file. Please upload a .csv or .xlsx file.'


EXPORT_COLUMNS = ['section_id', 'section_name', 'question_number'] + CSV_TEMPLATE_COLUMNS


@admin_bp.route('/api/test-questions/export/<int:description_id>', methods=['GET'])
@admin_login_required
def export_questions(description_id):
    admin_email = session.get('admin_email', 'UnknownAdmin')
    fmt = (request.args.get('fmt') or 'xlsx').lower()
    try:
        desc = user_op.get_test_description_by_id(description_id)
        if not desc:
            logger.warning(f"EXPORT BLOCKED: [{admin_email}] description {description_id} not found")
            return jsonify({'message': 'Test Description not found.'}), 404

        questions = user_op.get_all_questions_for_description(description_id)
        logger.info(f"EXPORT 1: [{admin_email}] desc={description_id} fmt={fmt} rows={len(questions)}")

        def row_values(q):
            return [q.get('section_id') or '', q.get('section_name') or '', q.get('question_number') or '',
                    q.get('question_type') or '', q.get('question_text') or '',
                    q.get('option_a') or '', q.get('option_b') or '', q.get('option_c') or '', q.get('option_d') or '',
                    q.get('correct_option') or '', q.get('correct_marks') if q.get('correct_marks') is not None else '',
                    q.get('negative_marks') if q.get('negative_marks') is not None else '',
                    q.get('question_level') or '', q.get('answer_text') or '', q.get('answer_link') or '']

        safe_key = re.sub(r'[^A-Za-z0-9_-]', '', str(desc.get('test_key') or description_id))

        if fmt == 'csv':
            out = StringIO()
            writer = csv.writer(out)
            writer.writerow(EXPORT_COLUMNS)
            for q in questions:
                writer.writerow(row_values(q))
            logger.info(f"EXPORT 2: [{admin_email}] csv ready for desc={description_id}")
            return Response(out.getvalue(), mimetype='text/csv',
                            headers={'Content-Disposition': f'attachment; filename=questions_{safe_key}.csv'})

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            logger.error("EXPORT: openpyxl not installed; falling back unavailable")
            return jsonify({'message': 'Excel support is not installed. Use ?fmt=csv instead.'}), 501

        wb = Workbook()
        ws = wb.active
        ws.title = 'Questions'
        ws.append(EXPORT_COLUMNS)
        for q in questions:
            ws.append(row_values(q))
        header_fill = PatternFill('solid', fgColor='833D17')
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = header_fill
        for idx, width in enumerate([12, 16, 10, 12, 55, 20, 20, 20, 20, 12, 12, 12, 12, 40, 20], 1):
            ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        logger.info(f"EXPORT 2: [{admin_email}] xlsx ready for desc={description_id}")
        return send_file(buf, as_attachment=True, download_name=f'questions_{safe_key}.xlsx',
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    except Exception:
        logger.exception(f"EXPORT FAILED: [{admin_email}] desc={description_id}")
        return jsonify({'message': 'Could not export the questions.'}), 500


@admin_bp.route('/api/questions/<int:description_id>', methods=['GET'])
@admin_login_required
def get_questions_api(description_id):
    admin_email = session.get('admin_email', 'UnknownAdmin')
    try:
        section_id = request.args.get('section_id')
        if not section_id:
            return jsonify({'message': 'Section ID is required'}), 400
        desc = user_op.get_test_description_by_id(description_id)
        if not desc:
            return jsonify({'message': 'Test Description not found'}), 404
        questions = user_op.get_questions_for_description_section(description_id, section_id)
        current_count = user_op.get_question_count_for_description(description_id)
        max_allowed = desc.get('total_questions', 0)
        logger.info(f"[{admin_email}] Retrieved questions for description_id={description_id}, section_id={section_id}")
        return jsonify({
            'questions': questions,
            'current_count': current_count,
            'max_allowed': max_allowed
        })
    except Exception as e:
        logger.exception(f"[{admin_email}] Failed to fetch questions for description_id={description_id}")
        return jsonify({'message': 'Could not load questions. Please refresh the page.'}), 500

@admin_bp.route('/api/test-questions/add', methods=['POST'])
@admin_login_required
def add_test_question_api():
    admin_email = session.get('admin_email', 'UnknownAdmin')
    data = request.get_json()
    required = ['test_id', 'test_key', 'test_description_id', 'section_id', 'section_name',
                'question_type', 'question_text', 'correct_option', 'correct_marks']
    if not all(data.get(field) is not None for field in required):
        logger.warning(f"[{admin_email}] Missing required fields in test question add.")
        return jsonify({'message': 'Missing required fields.'}), 400
    try:
        logger.info(f"ADD Q 1: [{admin_email}] desc={data.get('test_description_id')} section={data.get('section_id')!r}")

        data['negative_marks'] = _convert_negative_marks(data.get('negative_marks'))
        data['correct_marks'] = float(data['correct_marks'])

        if not _has_text(data.get('question_text')) and '<img' not in str(data.get('question_text') or ''):
            return jsonify({'message': 'Write the question text.'}), 400
        qtype, answer = _check_answer(data.get('question_type'), data.get('correct_option'),
                                      {k: data.get(f'option_{k}') for k in 'abcd'})
        data['question_type'] = qtype
        data['correct_option'] = answer
        fp = question_fingerprint(qtype, data.get('question_text'), data.get('option_a'), data.get('option_b'),
                                  data.get('option_c'), data.get('option_d'))
        already = stored_fingerprints(data['test_description_id']).get(fp)
        if already:
            logger.warning(f"ADD Q BLOCKED: [{admin_email}] duplicate of {where_stored(already)} in description {data['test_description_id']}")
            return jsonify({'message': f'This exact question is already in this test as {where_stored(already)}. '
                                       f'Nothing was added - change the text or options if it is a different question.'}), 409

        desc = user_op.get_test_description_by_id(data['test_description_id'])
        if not desc:
            logger.warning(f"ADD Q 3 BLOCKED: [{admin_email}] description {data['test_description_id']} not found")
            return jsonify({'message': 'Test Description not found.'}), 404

        declared = int(desc.get('total_questions') or 0)
        current = user_op.get_question_count_for_description(data['test_description_id'])
        logger.info(f"ADD Q 4: [{admin_email}] current={current}/{declared} questions")

        if current >= declared:
            logger.warning(f"ADD Q 5 BLOCKED: [{admin_email}] limit reached ({current}/{declared})")
            return jsonify({
                'message': f"You set {declared} questions in the test description and all {current} are already uploaded. "
                           f"Delete a question or increase Total Questions in the description first."
            }), 409

        for attempt in range(3):
            last_q_num = user_op.get_max_question_number_in_section(data['test_description_id'], data['section_id'])
            data['question_number'] = last_q_num + 1
            try:
                user_op.insert_test_question(data)
                break
            except mysql.connector.Error as err:
                if err.errno != 1062 or attempt == 2:
                    raise
                logger.info(f"ADD Q: [{admin_email}] number {data['question_number']} was taken, retrying")

        new_count = current + 1
        remaining = declared - new_count
        declared_marks = float(desc.get('total_marks') or 0)
        marks_now = user_op.get_marks_sum_for_description(data['test_description_id'])
        logger.info(f"ADD Q 6: [{admin_email}] saved. {new_count}/{declared} questions, marks {marks_now}/{declared_marks}")

        message = f"Question added ({new_count}/{declared})."
        if remaining == 0:
            message += " The test is now full."
        else:
            message += f" {remaining} more can be added."
        if marks_now > declared_marks:
            message += (f" Warning: uploaded marks now total {marks_now:g}, which exceeds the "
                        f"{declared_marks:g} declared in the description.")

        return jsonify({'message': message, 'current_count': new_count,
                        'max_allowed': declared, 'remaining': remaining}), 201
    except ValueError as e:
        logger.warning(f"ADD Q BLOCKED (validation): [{admin_email}] {e}")
        return jsonify({'message': str(e)}), 400
    except Exception:
        logger.exception(f"ADD Q FAILED: [{admin_email}] could not add test question")
        return jsonify({'message': 'Could not save the question. Please try again.'}), 500

@admin_bp.route('/api/test-questions/update/<int:id>', methods=['PUT'])
@admin_login_required
def update_test_question_api(id):
    admin_email = session.get('admin_email', 'UnknownAdmin')
    if not can_manage():
        return jsonify({'message': 'Permission Denied.'}), 403
    data = request.get_json(silent=True)
    if not data:
        return jsonify({'message': 'No data provided.'}), 400
    current = user_op.get_question(id)
    if not current:
        return jsonify({'message': 'Question not found. Refresh the list.'}), 404
    try:
        if 'negative_marks' in data:
            data['negative_marks'] = _convert_negative_marks(data['negative_marks'])
        if 'correct_marks' in data:
            data['correct_marks'] = float(data['correct_marks'])
            if data['correct_marks'] <= 0:
                raise ValueError('Correct marks must be more than 0.')
        merged = dict(current, **data)
        if not current.get('question_image'):
            if not _has_text(merged.get('question_text')) and '<img' not in str(merged.get('question_text') or ''):
                raise ValueError('Write the question text.')
            qtype, answer = _check_answer(merged.get('question_type'), merged.get('correct_option'),
                                          {k: merged.get(f'option_{k}') for k in 'abcd'})
            data['question_type'], data['correct_option'] = qtype, answer
        user_op.update_test_question_fields(id, data)
        logger.info(f"[{admin_email}] Updated test question ID {id}: {data}")
        return jsonify({'message': 'Question updated successfully!'})
    except ValueError as e:
        logger.warning(f"[{admin_email}] Validation error: {e}")
        return jsonify({'message': str(e)}), 400
    except Exception as e:
        logger.exception(f"[{admin_email}] Failed to update test question ID {id}")
        return jsonify({'message': 'Could not update the question. Please try again.'}), 500

@admin_bp.route('/api/test-questions/delete/<int:id>', methods=['DELETE'])
@admin_login_required
def delete_test_question_api(id):
    if not can_manage():
        logger.warning("DELETE Q BLOCKED: [%s] role may not manage questions",
                       session.get('admin_email', 'UnknownAdmin'))
        return jsonify({'message': 'Permission Denied.'}), 403
    admin_email = session.get('admin_email', 'UnknownAdmin')
    try:
        row = user_op.get_question(id)
        if not row:
            return jsonify({'message': 'Question not found. Refresh the list.'}), 404
        user_op.delete_test_question(id)
        if int(id) >= COMM_ID_OFFSET:
            user_op.renumber_comm_section(row['test_key'], row['section_name'])
        else:
            user_op.renumber_section(row['test_description_id'], row['section_id'])
        logger.info(f"[{admin_email}] Deleted test question ID {id}; section renumbered")
        return jsonify({'message': 'Question deleted successfully!'})
    except Exception as e:
        logger.exception(f"[{admin_email}] Failed to delete test question ID {id}")
        return jsonify({'message': 'Could not delete the question. Please try again.'}), 500

@admin_bp.route('/api/test-questions/upload-csv', methods=['POST'])
@admin_login_required
def upload_questions_csv():
    if 'csv_file' not in request.files: return jsonify({"message": "No file provided"}), 400
    csv_file = request.files['csv_file']
    if not csv_file or not csv_file.filename or not csv_file.filename.lower().endswith(('.csv', '.xlsx')):
        return jsonify({"message": "Invalid file. Please upload a .csv or .xlsx file."}), 400
    test_id = request.form.get('test_id')
    test_key = request.form.get('test_key')
    description_id = request.form.get('test_description_id')
    section_id = request.form.get('section_id')
    section_name = request.form.get('section_name')
    if not all([test_id, test_key, description_id, section_id, section_name]):
        return jsonify({"message": "Missing test/section identifiers."}), 400
    admin_email = session.get('admin_email', 'UnknownAdmin')
    try:
        desc = user_op.get_test_description_by_id(description_id)
        if not desc:
            logger.warning(f"CSV UPLOAD BLOCKED: [{admin_email}] description {description_id} not found")
            return jsonify({"message": "Test Description not found."}), 404

        declared = int(desc.get('total_questions') or 0)
        current = user_op.get_question_count_for_description(description_id)
        remaining = max(0, declared - current)
        logger.info(f"CSV UPLOAD 1: [{admin_email}] desc={description_id} current={current}/{declared}, room for {remaining}")

        if remaining == 0:
            logger.warning(f"CSV UPLOAD BLOCKED: [{admin_email}] test already full ({current}/{declared})")
            return jsonify({
                "message": f"You set {declared} questions in the test description and all {current} are already uploaded. "
                           f"No rows were imported. Delete questions or increase the limit first."
            }), 409

        reader, parse_error = _rows_from_upload(csv_file)
        if parse_error:
            logger.warning(f"CSV UPLOAD BLOCKED: [{admin_email}] {parse_error}")
            return jsonify({"message": parse_error}), 400
        logger.info(f"CSV UPLOAD 1b: [{admin_email}] parsed {len(reader)} data row(s) from {csv_file.filename!r}")

        valid_rows, errors = [], []
        last_q_num = user_op.get_max_question_number_in_section(description_id, section_id)
        cell = lambda row, key, default='': str(row.get(key) if row.get(key) is not None else default).strip()
        for i, row in enumerate(reader, 2):
            if not any(str(v or '').strip() for v in row.values()):
                continue
            try:
                if not cell(row, 'question_text'):
                    raise ValueError('question_text is empty.')
                options = {k: cell(row, f'option_{k}') for k in 'abcd'}
                qtype, answer = _check_answer(cell(row, 'question_type'), cell(row, 'correct_option'),
                                              None if cell(row, 'question_type').upper() == 'NAT' else options)
                try:
                    marks = float(cell(row, 'correct_marks', '0') or 0)
                except ValueError:
                    raise ValueError(f"correct_marks must be a number (got '{cell(row, 'correct_marks')}').")
                if marks <= 0:
                    raise ValueError('correct_marks must be more than 0.')
                level = (cell(row, 'question_level') or 'Medium').capitalize()
                if level not in ('Easy', 'Medium', 'Tough'):
                    raise ValueError(f"question_level must be Easy, Medium or Tough (got '{cell(row, 'question_level')}').")
                valid_rows.append({
                    'row_number': i,
                    'test_id': test_id,
                    'test_key': test_key,
                    'test_description_id': description_id,
                    'section_id': section_id,
                    'section_name': section_name,
                    'question_number': last_q_num + len(valid_rows) + 1,
                    'question_type': qtype,
                    'question_text': cell(row, 'question_text'),
                    'option_a': options['a'] or None,
                    'option_b': options['b'] or None,
                    'option_c': options['c'] or None,
                    'option_d': options['d'] or None,
                    'correct_option': answer,
                    'correct_marks': marks,
                    'negative_marks': _convert_negative_marks(cell(row, 'negative_marks', '0')),
                    'question_level': level,
                    'answer_text': cell(row, 'answer_text') or None,
                    'answer_link': cell(row, 'answer_link') or None,
                })
            except ValueError as e:
                errors.append({"row_number": i, "message": str(e)})

        if not valid_rows and not errors:
            return jsonify({"message": "The file has no question rows."}), 400
        stored = stored_fingerprints(description_id)
        seen, fresh, skipped = {}, [], []
        for row in valid_rows:
            fp = question_fingerprint(row['question_type'], row['question_text'], row['option_a'], row['option_b'], row['option_c'], row['option_d'])
            if fp in stored:
                skipped.append({"row_number": row['row_number'], "message": f"Already in this test as {where_stored(stored[fp])} - skipped."})
            elif fp in seen:
                skipped.append({"row_number": row['row_number'], "message": f"Same as row {seen[fp]} of this file - skipped."})
            else:
                seen[fp] = row['row_number']
                row['question_number'] = last_q_num + len(fresh) + 1
                fresh.append(row)
        for row in fresh:
            row.pop('row_number', None)
        valid_rows = fresh
        total_rows = len(valid_rows) + len(errors)
        if total_rows > remaining:
            errors.append({"row_number": '-',
                           "message": f"The file has {total_rows} questions but only {remaining} more fit in this test "
                                      f"({current}/{declared} used). Raise Total Questions or remove rows."})
        if errors:
            logger.warning(f"CSV UPLOAD REFUSED: [{admin_email}] {len(errors)} problem(s), nothing saved")
            return jsonify({
                "message": f"Nothing was saved - fix the {len(errors)} problem(s) below and upload the same file again.",
                "inserted_count": 0, "failed_count": len(errors), "errors": errors,
                "current_count": current, "max_allowed": declared
            }), 400

        failed = 0
        if not valid_rows:
            logger.info(f"CSV UPLOAD: [{admin_email}] every row already stored ({len(skipped)} skipped), nothing added")
            return jsonify({"message": f"Nothing added - all {len(skipped)} question(s) in this file are already in this test.",
                            "inserted_count": 0, "failed_count": 0, "errors": [], "skipped_count": len(skipped), "skipped": skipped,
                            "current_count": current, "max_allowed": declared}), 200
        try:
            inserted = user_op.insert_test_questions_bulk(valid_rows)
            logger.info(f"CSV UPLOAD 1c: [{admin_email}] bulk-inserted {inserted} row(s) atomically, {len(skipped)} duplicate(s) skipped")
        except Exception as exc:
            logger.exception(f"CSV UPLOAD FAILED: [{admin_email}] bulk insert rolled back")
            from admin.db_errors import friendly_db_error
            exact = friendly_db_error(exc)
            return jsonify({
                "message": (exact or "The file was read but saving failed.")
                           + " NO questions were imported - fix the issue and upload the same file again.",
                "inserted_count": 0, "failed_count": len(valid_rows), "errors": []
            }), 500

        new_count = current + inserted
        marks_now = user_op.get_marks_sum_for_description(description_id)
        declared_marks = float(desc.get('total_marks') or 0)
        message = f"{inserted} new question(s) added. Test now has {new_count}/{declared} questions."
        if skipped:
            message += f" {len(skipped)} already in this test were skipped."
        if marks_now > declared_marks:
            message += (f" Warning: uploaded marks total {marks_now:g}, which exceeds the "
                        f"{declared_marks:g} declared in the description.")
        logger.info(f"CSV UPLOAD 2: [{admin_email}] inserted={inserted} failed={failed} -> {new_count}/{declared}")
        return jsonify({
            "message": message, "inserted_count": inserted,
            "failed_count": failed, "errors": errors, "skipped_count": len(skipped), "skipped": skipped,
            "current_count": new_count, "max_allowed": declared
        }), 200
    except Exception as e:
        logger.exception(f"CSV UPLOAD FAILED: [{admin_email}] {e}")
        return jsonify({'message': 'Failed to process the CSV file. Please check the format and try again.'}), 500

@admin_bp.route('/api/test-sections/<int:description_id>', methods=['GET'])
@admin_login_required
def list_test_sections_api(description_id):
    desc = user_op.get_test_description_by_id(description_id)
    if not desc:
        return jsonify({'message': 'Test description not found.'}), 404
    rows = user_op.sections_for_description(description_id)
    return jsonify({'sections': [{'section_id': r['section_id'], 'section_name': r['section_name'],
                                  'questions': int(r['questions'] or 0), 'marks': float(r['marks'] or 0)}
                                 for r in rows],
                    'current_count': user_op.get_question_count_for_description(description_id),
                    'max_allowed': int(desc.get('total_questions') or 0)})


@admin_bp.route('/api/test-questions/bulk-delete', methods=['POST'])
@admin_login_required
def bulk_delete_questions_api():
    admin_email = session.get('admin_email', 'UnknownAdmin')
    if not can_manage():
        return jsonify({'message': 'Permission Denied.'}), 403
    data = request.get_json(silent=True) or {}
    try:
        description_id = int(data.get('description_id'))
        ids = sorted({int(i) for i in (data.get('ids') or [])})
    except (TypeError, ValueError):
        return jsonify({'message': 'Choose the questions to delete.'}), 400
    if not ids:
        return jsonify({'message': 'Choose at least one question.'}), 400
    if len(ids) > 2000:
        return jsonify({'message': 'Delete at most 2000 questions at a time.'}), 400
    if not user_op.get_test_description_by_id(description_id):
        return jsonify({'message': 'Test description not found.'}), 404
    found = user_op.questions_in_description(description_id, ids)
    if len(found) != len(ids):
        return jsonify({'message': f'{len(ids) - len(found)} of the chosen questions are no longer in this test. '
                                   f'Refresh the list and try again.'}), 409
    used = user_op.answers_for_questions(description_id, ids)
    if data.get('dry_run'):
        return jsonify({'count': len(ids), **used})
    try:
        deleted = user_op.delete_questions(description_id, ids)
        if description_id >= COMM_ID_OFFSET:
            desc = user_op.get_test_description_by_id(description_id)
            for name in {r['section_name'] for r in found}:
                user_op.renumber_comm_section(desc['test_key'], name)
        else:
            for sid in {r['section_id'] for r in found}:
                user_op.renumber_section(description_id, sid)
    except Exception:
        logger.exception(f"BULK DELETE FAILED: [{admin_email}] description {description_id}")
        return jsonify({'message': 'Could not delete the questions. Nothing was removed.'}), 500
    logger.warning(f"BULK DELETE: [{admin_email}] removed {deleted} question(s) from description {description_id} "
                   f"({used['answers']} student answer(s) from {used['students']} student(s) went with them)")
    return jsonify({'message': f'{deleted} question(s) deleted.', 'deleted': deleted,
                    'current_count': user_op.get_question_count_for_description(description_id)})


def _section_id_ok(value):
    return bool(re.match(r'^[A-Za-z0-9_-]{1,50}$', value or ''))


def _chosen(data):
    description_id = int(data.get('description_id'))
    ids = sorted({int(i) for i in (data.get('ids') or [])})
    if not ids:
        raise ValueError('Choose at least one question.')
    if len(ids) > 2000:
        raise ValueError('Choose at most 2000 questions at a time.')
    if not user_op.get_test_description_by_id(description_id):
        raise LookupError('Test description not found.')
    found = user_op.questions_in_description(description_id, ids)
    if len(found) != len(ids):
        raise PermissionError(f'{len(ids) - len(found)} of the chosen questions are not in this test. Refresh the list and try again.')
    return description_id, ids, found


def _chosen_or_error(data):
    try:
        return _chosen(data), None
    except (TypeError, ValueError) as e:
        return None, (jsonify({'message': str(e) or 'Choose the questions first.'}), 400)
    except LookupError as e:
        return None, (jsonify({'message': str(e)}), 404)
    except PermissionError as e:
        return None, (jsonify({'message': str(e)}), 409)


@admin_bp.route('/api/test-questions/bulk-marks', methods=['POST'])
@admin_login_required
def bulk_marks_api():
    admin_email = session.get('admin_email', 'UnknownAdmin')
    if not can_manage():
        return jsonify({'message': 'Permission Denied.'}), 403
    data = request.get_json(silent=True) or {}
    chosen, err = _chosen_or_error(data)
    if err:
        return err
    description_id, ids, _ = chosen
    correct = negative = None
    try:
        if str(data.get('correct_marks', '')).strip() != '':
            correct = float(data['correct_marks'])
            if correct <= 0:
                return jsonify({'message': 'Correct marks must be more than 0.'}), 400
        if str(data.get('negative_marks', '')).strip() != '':
            negative = _convert_negative_marks(data['negative_marks'])
            if negative < 0:
                return jsonify({'message': 'Negative marks cannot be below 0.'}), 400
    except ValueError as e:
        return jsonify({'message': str(e)}), 400
    if correct is None and negative is None:
        return jsonify({'message': 'Enter correct marks, negative marks, or both.'}), 400
    try:
        n = user_op.update_marks(description_id, ids, correct, negative)
    except Exception:
        logger.exception(f"BULK MARKS FAILED: [{admin_email}] description {description_id}")
        return jsonify({'message': 'Could not update the marks. Nothing was changed.'}), 500
    logger.info(f"BULK MARKS: [{admin_email}] {n} question(s) in description {description_id} -> +{correct} / -{negative}")
    bits = []
    if correct is not None:
        bits.append(f'correct marks {correct:g}')
    if negative is not None:
        bits.append(f'negative marks {negative:g}')
    return jsonify({'message': f"{n} question(s) updated: {' and '.join(bits)}.", 'updated': n})


@admin_bp.route('/api/test-questions/bulk-move', methods=['POST'])
@admin_login_required
def bulk_move_api():
    admin_email = session.get('admin_email', 'UnknownAdmin')
    if not can_manage():
        return jsonify({'message': 'Permission Denied.'}), 403
    data = request.get_json(silent=True) or {}
    chosen, err = _chosen_or_error(data)
    if err:
        return err
    description_id, ids, found = chosen
    try:
        if description_id >= COMM_ID_OFFSET:
            desc = user_op.get_test_description_by_id(description_id)
            name = (data.get('section_name') or '').strip()
            target = next((sec for sec in user_op.get_comm_sections(desc['test_key']) if sec['section_name'] == name), None)
            if not target:
                return jsonify({'message': 'Choose an existing communication section (step 1).'}), 400
            if all(r['section_name'] == name for r in found):
                return jsonify({'message': f'Those questions are already in {name}.'}), 400
            n = user_op.move_comm_questions(desc['test_key'], ids, target)
            label = name
        else:
            target_id = (data.get('section_id') or '').strip()
            target_name = (data.get('section_name') or '').strip()
            if not _section_id_ok(target_id):
                return jsonify({'message': 'Section ID: 1-50 letters, numbers, - or _ (no spaces).'}), 400
            if not target_name:
                return jsonify({'message': 'Give the section a name.'}), 400
            existing = next((sec for sec in user_op.sections_for_description(description_id)
                             if sec['section_id'].lower() == target_id.lower()), None)
            if existing:
                target_id, target_name = existing['section_id'], existing['section_name']
            if all(r['section_id'] == target_id for r in found):
                return jsonify({'message': f'Those questions are already in {target_name} ({target_id}).'}), 400
            n = user_op.move_questions(description_id, ids, target_id, target_name, [r['section_id'] for r in found])
            label = f'{target_name} ({target_id})'
    except Exception:
        logger.exception(f"BULK MOVE FAILED: [{admin_email}] description {description_id}")
        return jsonify({'message': 'Could not move the questions. Nothing was changed.'}), 500
    logger.info(f"BULK MOVE: [{admin_email}] {n} question(s) in description {description_id} -> {label}")
    return jsonify({'message': f'{n} question(s) moved to {label}. Both sections were renumbered.', 'moved': n})


@admin_bp.route('/api/test-sections/<int:description_id>/rename', methods=['POST'])
@admin_login_required
def rename_section_api(description_id):
    admin_email = session.get('admin_email', 'UnknownAdmin')
    if not can_manage():
        return jsonify({'message': 'Permission Denied.'}), 403
    if description_id >= COMM_ID_OFFSET:
        return jsonify({'message': 'Communication sections are renamed in their Sections table.'}), 400
    if not user_op.get_test_description_by_id(description_id):
        return jsonify({'message': 'Test description not found.'}), 404
    data = request.get_json(silent=True) or {}
    section_id = (data.get('section_id') or '').strip()
    new_name = (data.get('section_name') or '').strip()[:100]
    if not new_name:
        return jsonify({'message': 'Give the section a name.'}), 400
    current = next((sec for sec in user_op.sections_for_description(description_id) if sec['section_id'] == section_id), None)
    if not current:
        return jsonify({'message': 'Section not found. Refresh the page.'}), 404
    if current['section_name'] == new_name:
        return jsonify({'message': 'That is already the section name.', 'renamed': 0})
    try:
        n = user_op.rename_section(description_id, section_id, new_name)
    except Exception:
        logger.exception(f"RENAME SECTION FAILED: [{admin_email}] {description_id}/{section_id}")
        return jsonify({'message': 'Could not rename the section.'}), 500
    logger.info(f"RENAME SECTION: [{admin_email}] {description_id}/{section_id} '{current['section_name']}' -> '{new_name}' ({n} question(s))")
    return jsonify({'message': f'Section {section_id} renamed to "{new_name}" on {n} question(s).', 'renamed': n})


@admin_bp.route('/api/test-descriptions/<int:id>/tcs-publish', methods=['POST'])
@admin_login_required
def tcs_publish_api(id):
    admin_email = session.get('admin_email', 'UnknownAdmin')
    if not can_manage():
        return jsonify({'message': 'Permission Denied.'}), 403
    desc = user_op.get_test_description_by_id(id)
    if not desc or user_op.section_for_test_id(desc.get('test_id')) != 'tcs_test':
        return jsonify({'message': 'This is not a TCS exam.'}), 404
    from admin.manage_mock_test.coding_question_panel import ensure_tcs_exam
    exam, cp_test = ensure_tcs_exam(desc)
    want = bool((request.get_json(silent=True) or {}).get('published'))
    summary = user_op.tcs_exam_summary(desc['test_key'])
    if want:
        problems = []
        if summary['aptitude_questions'] < int(desc.get('total_questions') or 0):
            problems.append(f"aptitude has {summary['aptitude_questions']} of {desc.get('total_questions')} questions")
        if summary['coding_questions'] == 0:
            problems.append('the coding part has no problem yet')
        if problems:
            return jsonify({'message': 'Cannot publish yet: ' + ' and '.join(problems) + '.'}), 400
    user_op._execute_query("UPDATE tcs_exams SET is_published = %s WHERE id = %s", (int(want), exam['id']))
    user_op._execute_query("UPDATE cp_tests SET is_published = 1 WHERE id = %s", (cp_test['id'],))
    logger.info(f"TCS EXAM: [{admin_email}] {'published' if want else 'unpublished'} exam {exam['id']} ({desc['test_key']})")
    return jsonify({'message': f"\"{desc['subject_title']}\" is now {'visible to students' if want else 'hidden from students (draft)'}.",
                    'published': want})
