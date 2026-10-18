import csv
import mysql.connector
import logging
from io import StringIO, BytesIO
from flask import request, jsonify, session, Response, send_file
from admin.manage_mock_test.test_details_db import UserOperation, COMM_ID_OFFSET
from extensions import admin_login_required, can_manage
from admin import admin_bp

logger = logging.getLogger(__name__)
user_op = UserOperation()

COMM_TYPES = ('MCQ', 'SPEAK', 'THINK_SPEAK', 'WRITING', 'COMPREHENSION', 'LISTEN_SPEAK', 'FILL_BLANK')


def _whole(value, name):
    text = str(value if value is not None else '').strip()
    if not text:
        return 0
    try:
        number = float(text)
    except ValueError:
        raise ValueError(f'{name} must be a whole number (got "{text}").')
    if number < 0 or number != int(number):
        raise ValueError(f'{name} must be a whole number of 0 or more (got "{text}").')
    return int(number)


def _validate_comm_question(row, sections_by_name):
    qtype = str(row.get('question_type') or '').strip().upper()
    if qtype not in COMM_TYPES:
        raise ValueError(f"question_type must be one of {', '.join(COMM_TYPES)} (got '{qtype or 'empty'}').")
    section_name = str(row.get('section_name') or '').strip()
    if section_name not in sections_by_name:
        raise ValueError(f"section_name '{section_name}' does not exist - create the section first.")
    question_text = str(row.get('question_text') or '').strip()
    if not question_text:
        raise ValueError('question_text is required.')

    q = {
        'section_name': section_name,
        'section_id': sections_by_name[section_name]['section_id'],
        'question_type': qtype,
        'instruction_text': (row.get('instruction_text') or '').strip() or None,
        'question_text': question_text,
        'media_url': (row.get('media_url') or '').strip() or None,
        'options_title': (row.get('options_title') or '').strip() or None,
        'option_a': (row.get('option_a') or '').strip() or None,
        'option_b': (row.get('option_b') or '').strip() or None,
        'option_c': (row.get('option_c') or '').strip() or None,
        'option_d': (row.get('option_d') or '').strip() or None,
        'correct_option': (str(row.get('correct_option') or '').strip().lower() or None),
        'recording_time_seconds': _whole(row.get('recording_time_seconds'), 'recording_time_seconds'),
        'thinking_time_seconds': _whole(row.get('thinking_time_seconds'), 'thinking_time_seconds'),
        'word_limit': _whole(row.get('word_limit'), 'word_limit'),
    }

    if qtype in ('MCQ', 'COMPREHENSION'):
        if not (q['option_a'] and q['option_b']):
            raise ValueError(f'{qtype} needs at least option_a and option_b.')
        if not q['correct_option'] or q['correct_option'] not in ('a', 'b', 'c', 'd'):
            raise ValueError(f'{qtype} needs correct_option a/b/c/d.')
        if not q[f"option_{q['correct_option']}"]:
            raise ValueError(f"correct_option is '{q['correct_option']}' but that option is empty.")
    if qtype in ('SPEAK', 'LISTEN_SPEAK') and q['recording_time_seconds'] <= 0:
        q['recording_time_seconds'] = 45
    if qtype == 'THINK_SPEAK':
        if q['thinking_time_seconds'] <= 0:
            q['thinking_time_seconds'] = 30
        if q['recording_time_seconds'] <= 0:
            q['recording_time_seconds'] = 60
    if qtype == 'WRITING' and q['word_limit'] <= 0:
        q['word_limit'] = 150
    return q


def _comm_context_or_error(description_id):
    if int(description_id) < COMM_ID_OFFSET:
        return None, (jsonify({'message': 'This is not a communication test description.'}), 400)
    desc = user_op.get_test_description_by_id(description_id)
    if not desc:
        return None, (jsonify({'message': 'Test description not found.'}), 404)
    section = user_op.section_for_test_id(desc.get('test_id'))
    if section == 'coding_test':
        return None, (jsonify({'message': 'This is a Coding Test description - '
                                          'use the coding question editor for it.'}), 400)
    return desc, None


@admin_bp.route('/api/comm-sections/<int:description_id>', methods=['GET'])
@admin_login_required
def list_comm_sections(description_id):
    desc, err = _comm_context_or_error(description_id)
    if err:
        return err
    return jsonify(user_op.get_comm_sections(desc['test_key']))


@admin_bp.route('/api/comm-sections/<int:description_id>', methods=['POST'])
@admin_login_required
def add_comm_section_api(description_id):
    admin_email = session.get('admin_email', 'UnknownAdmin')
    desc, err = _comm_context_or_error(description_id)
    if err:
        return err
    data = request.get_json(silent=True) or {}
    title = (data.get('section_title') or '').strip()
    if not title:
        return jsonify({'message': 'Section title is required.'}), 400
    try:
        duration = int(data.get('section_duration_seconds') or 0)
    except (TypeError, ValueError):
        return jsonify({'message': 'Section duration must be a number of seconds.'}), 400
    if duration <= 0:
        return jsonify({'message': 'Section duration (seconds) must be greater than zero.'}), 400

    existing = user_op.get_comm_sections(desc['test_key'])
    order = len(existing) + 1
    letter = chr(ord('A') + len(existing)) if len(existing) < 26 else str(len(existing) + 1)
    payload = {
        'test_key': desc['test_key'],
        'section_id': f'SEC_{letter}',
        'section_name': f'SECTION_{letter}',
        'section_title': title,
        'section_instruction': (data.get('section_instruction') or '').strip(),
        'section_duration_seconds': duration,
        'section_order': order,
    }
    user_op.add_comm_section(payload)
    logger.info(f"COMM SEC: [{admin_email}] added {payload['section_name']} '{title}' to {desc['test_key']}")
    return jsonify({'message': f"Section {payload['section_name']} - {title} added.", 'section': payload}), 201


@admin_bp.route('/api/comm-sections/update/<int:section_pk>', methods=['PUT'])
@admin_login_required
def update_comm_section_api(section_pk):
    admin_email = session.get('admin_email', 'UnknownAdmin')
    if not can_manage():
        return jsonify({'message': 'Permission Denied.'}), 403
    row = user_op._fetch_one("SELECT * FROM comm_sections WHERE id = %s", (section_pk,))
    if not row:
        return jsonify({'message': 'Section not found. Refresh the page.'}), 404
    data = request.get_json(silent=True) or {}
    title = (data.get('section_title') or '').strip()[:255]
    if not title:
        return jsonify({'message': 'Section title is required.'}), 400
    try:
        duration = int(data.get('section_duration_seconds') or 0)
    except (TypeError, ValueError):
        return jsonify({'message': 'Section duration must be a number of seconds.'}), 400
    if duration <= 0:
        return jsonify({'message': 'Section duration (seconds) must be greater than zero.'}), 400
    instruction = (data.get('section_instruction') or '').strip()
    try:
        user_op.update_comm_section(section_pk, title, instruction, duration)
    except Exception:
        logger.exception(f"COMM SEC UPDATE FAILED: [{admin_email}] {section_pk}")
        return jsonify({'message': 'Could not save the section.'}), 500
    logger.info(f"COMM SEC: [{admin_email}] updated section {row['section_name']} -> '{title}', {duration}s")
    return jsonify({'message': f"Section {row['section_name']} updated."})


@admin_bp.route('/api/comm-sections/delete/<int:section_pk>', methods=['DELETE'])
@admin_login_required
def delete_comm_section_api(section_pk):
    admin_email = session.get('admin_email', 'UnknownAdmin')
    n = user_op.delete_comm_section(section_pk)
    if not n:
        return jsonify({'message': 'Section not found.'}), 404
    logger.warning(f"COMM SEC: [{admin_email}] deleted section id={section_pk} and its questions")
    return jsonify({'message': 'Section and its questions deleted.'})


@admin_bp.route('/api/comm-questions/add/<int:description_id>', methods=['POST'])
@admin_login_required
def add_comm_question_api(description_id):
    admin_email = session.get('admin_email', 'UnknownAdmin')
    desc, err = _comm_context_or_error(description_id)
    if err:
        return err
    data = request.get_json(silent=True) or {}
    sections = {s['section_name']: s for s in user_op.get_comm_sections(desc['test_key'])}
    if not sections:
        return jsonify({'message': 'Create at least one section first.'}), 400

    declared = int(desc.get('total_questions') or 0)
    current = user_op.get_question_count_for_description(description_id)
    if current >= declared:
        return jsonify({'message': f"All {declared} questions are already uploaded for this test."}), 409

    try:
        q = _validate_comm_question(data, sections)
        from admin.manage_mock_test.test_details_insert import _convert_negative_marks
        negative = _convert_negative_marks(data.get('negative_marks'))
        marks = float(data.get('correct_marks') or 0)
    except (ValueError, TypeError) as e:
        return jsonify({'message': str(e) or 'Marks must be numbers.'}), 400

    from admin.manage_mock_test.test_details_insert import question_fingerprint, stored_fingerprints, where_stored
    already = stored_fingerprints(description_id).get(question_fingerprint(
        q['question_type'], q['question_text'], q.get('option_a'), q.get('option_b'), q.get('option_c'), q.get('option_d')))
    if already:
        return jsonify({'message': f'This exact question is already in this test as {where_stored(already)}. Nothing was added.'}), 409

    q.update({
        'test_id': desc['test_id'],
        'test_key': desc['test_key'],
        'test_description_id': description_id,
        'question_number': user_op.get_comm_max_question_number(desc['test_key'], q['section_name']) + 1,
        'correct_marks': marks,
        'negative_marks': negative,
        'question_level': data.get('question_level') or 'Medium',
    })
    try:
        for attempt in range(3):
            q['question_number'] = user_op.get_comm_max_question_number(desc['test_key'], q['section_name']) + 1
            try:
                user_op.insert_comm_question(q)
                break
            except mysql.connector.Error as err:
                if err.errno != 1062 or attempt == 2:
                    raise
    except Exception:
        logger.exception(f"COMM Q FAILED: [{admin_email}]")
        return jsonify({'message': 'Could not save the question. Please try again.'}), 500
    new_count = current + 1
    logger.info(f"COMM Q: [{admin_email}] {q['question_type']} question saved "
                f"({new_count}/{declared}) in {q['section_name']} of {desc['test_key']}")
    return jsonify({'message': f"{q['question_type']} question added ({new_count}/{declared}).",
                    'current_count': new_count, 'max_allowed': declared}), 201


@admin_bp.route('/api/comm-questions/update/<int:qid>', methods=['PUT'])
@admin_login_required
def update_comm_question_api(qid):
    admin_email = session.get('admin_email', 'UnknownAdmin')
    if not can_manage():
        return jsonify({'message': 'Permission Denied.'}), 403
    current = user_op.get_comm_question(qid)
    if not current:
        return jsonify({'message': 'Question not found. Refresh the list.'}), 404
    desc, err = _comm_context_or_error(current['test_description_id'])
    if err:
        return err
    data = request.get_json(silent=True) or {}
    sections = {s['section_name']: s for s in user_op.get_comm_sections(desc['test_key'])}
    try:
        q = _validate_comm_question(data, sections)
        from admin.manage_mock_test.test_details_insert import _convert_negative_marks
        q['negative_marks'] = _convert_negative_marks(data.get('negative_marks'))
        q['correct_marks'] = float(data.get('correct_marks') or 0)
    except ValueError as e:
        return jsonify({'message': str(e)}), 400
    q['question_level'] = data.get('question_level') or 'Medium'
    q['question_number'] = (current['question_number'] if q['section_name'] == current['section_name']
                            else user_op.get_comm_max_question_number(desc['test_key'], q['section_name']) + 1)
    try:
        user_op.update_comm_question(qid, q)
        if q['section_name'] != current['section_name']:
            user_op.renumber_comm_section(desc['test_key'], current['section_name'])
    except Exception:
        logger.exception(f"COMM Q UPDATE FAILED: [{admin_email}] {qid}")
        return jsonify({'message': 'Could not save the question. Please try again.'}), 500
    logger.info(f"COMM Q: [{admin_email}] updated question {qid} ({q['question_type']}) in {q['section_name']}")
    return jsonify({'message': f"Question updated ({q['question_type']}, {q['section_name']})."})


COMM_CSV_COLUMNS = ['section_name', 'question_type', 'instruction_text', 'question_text', 'media_url',
                    'options_title', 'option_a', 'option_b', 'option_c', 'option_d', 'correct_option',
                    'recording_time_seconds', 'thinking_time_seconds', 'word_limit',
                    'correct_marks', 'negative_marks', 'question_level']

COMM_CSV_ROWS = [
    ['SECTION_A', 'SPEAK', 'Read Aloud', 'The quick brown fox jumps over the lazy dog.', '', '', '', '', '', '', '', '45', '', '', '5', '0', 'Medium'],
    ['SECTION_B', 'THINK_SPEAK', 'Topic Speech', 'Describe your dream job and why it appeals to you.', '', '', '', '', '', '', '', '60', '30', '', '5', '0', 'Medium'],
    ['SECTION_C', 'MCQ', 'Grammar', 'She ___ to the market yesterday.', '', 'Select Option', 'go', 'went', 'gone', 'going', 'b', '', '', '', '2', '1/3', 'Easy'],
    ['SECTION_D', 'COMPREHENSION', 'Listen & Answer', 'What is the main topic of the conversation?', '/static/audio/sample_convo.mp3', 'Choose Answer:', 'Travel', 'Business', 'Weather', 'Food', 'b', '', '', '', '2', '0', 'Medium'],
    ['SECTION_E', 'WRITING', 'Email Writing', 'Write an email to your manager asking for sick leave.', '', '', '', '', '', '', '', '', '', '150', '10', '0', 'Tough'],
]


@admin_bp.route('/api/comm-questions/csv-template', methods=['GET'])
@admin_login_required
def comm_csv_template():
    fmt = (request.args.get('fmt') or 'csv').lower()
    if fmt == 'xlsx':
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            return jsonify({'message': 'Excel support is not installed. Use the CSV template.'}), 501
        wb = Workbook()
        ws = wb.active
        ws.title = 'Questions'
        ws.append(COMM_CSV_COLUMNS)
        for r in COMM_CSV_ROWS:
            ws.append(r)
        fill = PatternFill('solid', fgColor='833D17')
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = fill
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name='comm_question_template.xlsx',
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    out = StringIO()
    w = csv.writer(out)
    w.writerow(COMM_CSV_COLUMNS)
    for r in COMM_CSV_ROWS:
        w.writerow(r)
    return Response(out.getvalue(), mimetype='text/csv',
                    headers={'Content-Disposition': 'attachment; filename=comm_question_template.csv'})


@admin_bp.route('/api/comm-questions/upload-csv/<int:description_id>', methods=['POST'])
@admin_login_required
def upload_comm_questions_csv(description_id):
    admin_email = session.get('admin_email', 'UnknownAdmin')
    desc, err = _comm_context_or_error(description_id)
    if err:
        return err
    upload = request.files.get('csv_file')
    if not upload or not upload.filename:
        return jsonify({'message': 'No file provided.'}), 400

    from admin.manage_mock_test.test_details_insert import _rows_from_upload, _convert_negative_marks
    rows, parse_error = _rows_from_upload(upload)
    if parse_error:
        return jsonify({'message': parse_error}), 400

    sections = {s['section_name']: s for s in user_op.get_comm_sections(desc['test_key'])}
    if not sections:
        return jsonify({'message': 'Create the sections first, then upload questions.'}), 400

    declared = int(desc.get('total_questions') or 0)
    current = user_op.get_question_count_for_description(description_id)
    remaining = max(0, declared - current)
    if remaining == 0:
        return jsonify({'message': f"All {declared} questions are already uploaded. Nothing imported."}), 409

    prepared, errors = [], []
    counters = {}
    for i, row in enumerate(rows, 2):
        if not any(str(v or '').strip() for v in row.values()):
            continue
        try:
            q = _validate_comm_question(row, sections)
            level = str(row.get('question_level') or 'Medium').strip().capitalize()
            if level not in ('Easy', 'Medium', 'Tough'):
                raise ValueError(f'question_level must be Easy, Medium or Tough (got "{level}").')
            try:
                marks = float(str(row.get('correct_marks') or 0).strip() or 0)
            except ValueError:
                raise ValueError(f'correct_marks must be a number (got "{row.get("correct_marks")}").')
            q.update({'correct_marks': marks,
                      'negative_marks': _convert_negative_marks(row.get('negative_marks')),
                      'question_level': level})
            prepared.append((i, q))
        except (ValueError, TypeError) as e:
            errors.append({'row_number': i, 'message': str(e)})

    if not prepared and not errors:
        return jsonify({'message': 'The file has no question rows.'}), 400
    from admin.manage_mock_test.test_details_insert import question_fingerprint, stored_fingerprints, where_stored
    stored = stored_fingerprints(description_id)
    seen, fresh, skipped = {}, [], []
    for i, q in prepared:
        fp = question_fingerprint(q['question_type'], q['question_text'], q.get('option_a'), q.get('option_b'), q.get('option_c'), q.get('option_d'))
        if fp in stored:
            skipped.append({'row_number': i, 'message': f'Already in this test as {where_stored(stored[fp])} - skipped.'})
        elif fp in seen:
            skipped.append({'row_number': i, 'message': f'Same as row {seen[fp]} of this file - skipped.'})
        else:
            seen[fp] = i
            fresh.append((i, q))
    prepared = fresh
    if len(prepared) + len(errors) > remaining:
        errors.append({'row_number': '-', 'message': f'The file has {len(prepared) + len(errors)} questions but only '
                                                     f'{remaining} more fit in this test ({current}/{declared} used). '
                                                     f'Raise Total Questions or remove rows.'})
    if errors:
        logger.warning(f"COMM CSV: [{admin_email}] refused - {len(errors)} problem(s), nothing saved for {desc['test_key']}")
        return jsonify({'message': f'Nothing was saved - fix the {len(errors)} problem(s) below and upload the same file again.',
                        'inserted_count': 0, 'failed_count': len(errors), 'errors': errors,
                        'current_count': current, 'max_allowed': declared}), 400

    if not prepared:
        return jsonify({'message': f'Nothing added - all {len(skipped)} question(s) in this file are already in this test.',
                        'inserted_count': 0, 'failed_count': 0, 'errors': [], 'skipped_count': len(skipped), 'skipped': skipped,
                        'current_count': current, 'max_allowed': declared}), 200
    batch = []
    for _, q in prepared:
        sec = q['section_name']
        if sec not in counters:
            counters[sec] = user_op.get_comm_max_question_number(desc['test_key'], sec)
        counters[sec] += 1
        q.update({'test_id': desc['test_id'], 'test_key': desc['test_key'],
                  'test_description_id': description_id, 'question_number': counters[sec]})
        batch.append(q)
    try:
        inserted = user_op.insert_comm_questions_bulk(batch)
    except Exception:
        logger.exception(f"COMM CSV: [{admin_email}] bulk insert failed for {desc['test_key']}")
        return jsonify({'message': 'Could not save the questions. Nothing was saved - please try again.'}), 500

    new_count = current + inserted
    logger.info(f"COMM CSV: [{admin_email}] inserted={inserted} failed={len(errors)} -> {new_count}/{declared} for {desc['test_key']}")
    message = f'{inserted} new question(s) added. Test now has {new_count}/{declared} questions.'
    if skipped:
        message += f' {len(skipped)} already in this test were skipped.'
    return jsonify({
        'message': message,
        'inserted_count': inserted, 'failed_count': len(errors), 'errors': errors,
        'skipped_count': len(skipped), 'skipped': skipped,
        'current_count': new_count, 'max_allowed': declared
    })
